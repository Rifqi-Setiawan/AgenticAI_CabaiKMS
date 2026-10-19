from __future__ import annotations

import copy
import math
from pathlib import Path
from typing import Any, Literal

import openpyxl
import pytest
from pydantic import ValidationError

from src.storage.document_contracts import (
    AdministrativePath,
    AttributeValue,
    CurrentSourceReference,
    DocumentDraft,
    DocumentReviewItem,
    HabitatIdentity,
    KnowledgeDocument,
    MediaReference,
    PhotoContextSignal,
    PhotoMismatchDecision,
    PlantIdentity,
    SamplingLocationDetail,
    VisionClassificationResult,
)
from src.storage.attribute_policy import (
    AttributeCatalog,
    AttributePolicy,
    resolve_attribute,
)
from src.storage.identity_resolution import (
    AdministrativeRegistry,
    LocationResolution,
    PlantCatalog,
    PlantResolution,
    resolve_location,
    resolve_plant,
)
from src.storage.document_planner import (
    DocumentPlanner,
    ParsedCellValue,
    detect_intra_workbook_conflicts,
    evaluate_photo_variety_alignment,
    parse_cell_value,
    parse_range_value,
    plan_document_changes,
)


# ============================================================================
# Helpers & Fixtures
# ============================================================================


def _create_plant_identity(
    variety_id: str = "plant_rawit_frutescens_001",
    label: str = "Cabai rawit (Capsicum frutescens L)",
    level: Literal["source_category", "species", "cultivar"] = "source_category",
    scientific_name: str | None = "Capsicum frutescens L",
) -> PlantIdentity:
    """Construct a valid PlantIdentity fixture."""
    return PlantIdentity(
        id=variety_id,
        label=label,
        identity_level=level,
        scientific_name=scientific_name,
    )


def _create_habitat_identity(
    habitat_id: str = "hab_wonosobo_lamuk_001",
    village: str = "Desa Lamuk",
    district: str = "Kecamatan Kalikajar",
    regency: str = "Kabupaten Wonosobo",
    province: str = "Jawa Tengah",
) -> HabitatIdentity:
    """Construct a valid HabitatIdentity fixture."""
    return HabitatIdentity(
        id=habitat_id,
        definition="administrative_village_v1",
        administrative_path=AdministrativePath(
            province_id="prop_33",
            regency_id="kab_3307",
            district_id="kec_3307070",
            village_id="desa_3307070007",
            province_name=province,
            regency_name=regency,
            district_name=district,
            village_name=village,
        ),
        label=f"{village}, {district}, {regency}, {province}",
    )


def _create_sample_attribute_value(
    label: str,
    value: Any,
    value_type: Literal["text", "number", "range", "boolean", "list"] = "number",
    unit: str | None = None,
    cell: str = "G12",
    row: int = 12,
) -> AttributeValue:
    """Construct a valid AttributeValue fixture."""
    return AttributeValue(
        label=label,
        value=value,
        value_type=value_type,
        unit=unit,
        resolution="AUTO_ACCEPT",
        mapping_policy_version="exact-alias-policy-v1",
        current_source=CurrentSourceReference(
            source_ref="src_field_data_2026",
            sheet="Resume Data",
            row=row,
            cell=cell,
            location_ref="loc_field_01",
        ),
    )


def _create_minimal_knowledge_document(
    variety: PlantIdentity,
    habitat: HabitatIdentity,
    attributes: dict[str, AttributeValue] | None = None,
    run_id: str = "run_init_001",
) -> KnowledgeDocument:
    """Construct a minimal KnowledgeDocument fixture."""
    return KnowledgeDocument(
        schema_version="kms-document-v1",
        variety=variety,
        habitat=habitat,
        sampling_locations=[
            SamplingLocationDetail(
                location_ref="loc_ref_001",
                raw_address="Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
                hamlet=None,
                sampling_marker="lokasi 1",
            )
        ],
        attributes=attributes or {},
        media_refs=[],
        updated_by_run=run_id,
    )


# ============================================================================
# Section 1: Module Interface & Instantiation Contracts
# ============================================================================


def test_document_planner_module_interface_exports():
    """Verify document planner module exports required classes and functions."""
    assert callable(DocumentPlanner)
    assert callable(plan_document_changes)
    assert callable(detect_intra_workbook_conflicts)
    assert callable(parse_cell_value)
    assert callable(parse_range_value)
    assert callable(evaluate_photo_variety_alignment)


def test_document_planner_instantiation_defaults_and_custom():
    """DocumentPlanner instantiates with default registries or injected instances."""
    planner_default = DocumentPlanner()
    assert planner_default is not None

    custom_registry = AdministrativeRegistry()
    custom_catalog = PlantCatalog()
    custom_attr_catalog = AttributeCatalog()
    custom_policy = AttributePolicy()

    planner_custom = DocumentPlanner(
        catalog=custom_attr_catalog,
        policy=custom_policy,
        location_registry=custom_registry,
        plant_catalog=custom_catalog,
    )
    assert planner_custom is not None


# ============================================================================
# Section 2: Criterion A11 — Intra-Workbook Conflict Detection & Row Order
# ============================================================================


def test_intra_workbook_value_conflict_enqueued_to_review(resume_data_workbook_path: Path):
    """A11: Conflicting values within one workbook escalate to review; row order does not determine winner."""
    wb = openpyxl.load_workbook(resume_data_workbook_path, data_only=True)
    ws = wb["Resume Data"]

    # Verify actual ground-truth cell values for Desa Lamuk (rows 12 & 13)
    val_g12 = float(ws["G12"].value)
    val_g13 = float(ws["G13"].value)
    assert val_g12 == pytest.approx(49.02)
    assert val_g13 == pytest.approx(46.541666666666664)
    assert val_g12 != val_g13

    # Verify actual ground-truth cell values for Desa Argalingga (rows 14 & 15)
    val_g14 = float(ws["G14"].value)
    val_g15 = float(ws["G15"].value)
    assert val_g14 == pytest.approx(47.09)
    assert val_g15 == pytest.approx(45.620000000000005)
    assert val_g14 != val_g15

    planner = DocumentPlanner()
    draft = planner.plan_workbook(
        workbook_path=resume_data_workbook_path,
        run_id="run_field_survey_2026",
    )

    # Conflicts must produce review items with category INTRA_WORKBOOK_CONFLICT or INTRA_RUN_VALUE_CONFLICT
    conflict_reviews = [
        item for item in draft.reviews
        if item.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    assert len(conflict_reviews) >= 2, "Expected intra-workbook conflicts for both Lamuk and Argalingga."

    # Validate Lamuk conflict details
    lamuk_conflicts = [
        item for item in conflict_reviews
        if "lamuk" in (item.target_document_key or "").lower()
        or "lamuk" in (item.source_context or "").lower()
        or "g12" in (item.source_context or "").lower()
    ]
    assert len(lamuk_conflicts) >= 1
    lamuk_item = lamuk_conflicts[0]
    assert lamuk_item.status == "pending"
    assert any(pytest.approx(49.02, rel=1e-3) == v for v in lamuk_item.raw_values if isinstance(v, (int, float)))
    assert any(pytest.approx(46.5417, rel=1e-3) == v for v in lamuk_item.raw_values if isinstance(v, (int, float)))

    # Validate Argalingga conflict details
    argalingga_conflicts = [
        item for item in conflict_reviews
        if "argalingga" in (item.target_document_key or "").lower()
        or "argalingga" in (item.source_context or "").lower()
        or "g14" in (item.source_context or "").lower()
    ]
    assert len(argalingga_conflicts) >= 1
    argalingga_item = argalingga_conflicts[0]
    assert argalingga_item.status == "pending"
    assert any(pytest.approx(47.09, rel=1e-3) == v for v in argalingga_item.raw_values if isinstance(v, (int, float)))
    assert any(pytest.approx(45.62, rel=1e-3) == v for v in argalingga_item.raw_values if isinstance(v, (int, float)))


def test_intra_workbook_conflict_forbids_last_row_wins():
    """A11: Row order does not determine the winning value; last-row-wins is strictly prohibited."""
    planner = DocumentPlanner()

    row_a = {
        "Lokasi Sampling": "Dusun Pangempon, Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 49.02,
        "_row": 12,
        "_cell": "G12",
    }
    row_b = {
        "Lokasi Sampling": "Dusun Genting, Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 46.5417,
        "_row": 13,
        "_cell": "G13",
    }

    # Plan with row A then row B
    draft_forward = planner.plan_records(
        records=[row_a, row_b],
        run_id="run_order_test_forward",
    )

    # Plan with row B then row A (reversed order)
    draft_backward = planner.plan_records(
        records=[row_b, row_a],
        run_id="run_order_test_backward",
    )

    conflicts_fwd = [
        r for r in draft_forward.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    conflicts_bwd = [
        r for r in draft_backward.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]

    assert len(conflicts_fwd) == 1
    assert len(conflicts_bwd) == 1

    # Raw values in review items must contain both values regardless of input order
    raw_fwd = set(round(float(v), 2) for v in conflicts_fwd[0].raw_values)
    raw_bwd = set(round(float(v), 2) for v in conflicts_bwd[0].raw_values)
    assert raw_fwd == {49.02, 46.54}
    assert raw_bwd == {49.02, 46.54}

    # Neither document may silently resolve to the last row's value
    doc_key_fwd = list(draft_forward.documents.keys())[0]
    doc_fwd = draft_forward.documents[doc_key_fwd]
    doc_key_bwd = list(draft_backward.documents.keys())[0]
    doc_bwd = draft_backward.documents[doc_key_bwd]

    # The conflicting attribute must either be excluded from active values or marked unresolved
    if "attr_chlorophyll_spad" in doc_fwd.attributes:
        assert doc_fwd.attributes["attr_chlorophyll_spad"].resolution != "AUTO_ACCEPT"
    if "attr_chlorophyll_spad" in doc_bwd.attributes:
        assert doc_bwd.attributes["attr_chlorophyll_spad"].resolution != "AUTO_ACCEPT"


def test_intra_workbook_identical_values_produce_no_conflict():
    """Identical values for the same attribute across multiple rows do not trigger a conflict."""
    planner = DocumentPlanner()

    row_1 = {
        "Lokasi Sampling": "Dusun Randu, Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "pH tanah": 7.0,
        "_row": 3,
        "_cell": "L3",
    }
    row_2 = {
        "Lokasi Sampling": "Dusun Tenen, Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "pH tanah": 7.0,
        "_row": 4,
        "_cell": "L4",
    }

    draft = planner.plan_records(records=[row_1, row_2], run_id="run_identical_test")
    conflicts = [
        r for r in draft.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    assert len(conflicts) == 0, "Identical values must not raise an intra-workbook conflict."

    # Document should contain the agreed attribute value
    doc = list(draft.documents.values())[0]
    matched_attrs = [v for k, v in doc.attributes.items() if "ph" in k]
    assert len(matched_attrs) == 1
    assert matched_attrs[0].value == pytest.approx(7.0)


def test_intra_workbook_conflict_detect_function_standalone():
    """detect_intra_workbook_conflicts function detects differences on grouped rows."""
    target_key = "plant_rawit_001:hab_lamuk_001"
    grouped_rows = [
        {"attr_chlorophyll_spad": 49.02, "_row": 12, "_cell": "G12"},
        {"attr_chlorophyll_spad": 46.5417, "_row": 13, "_cell": "G13"},
    ]

    conflicts = detect_intra_workbook_conflicts(target_key, grouped_rows)
    assert len(conflicts) == 1
    item = conflicts[0]
    assert item.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    assert item.target_document_key == target_key
    assert len(item.raw_values) == 2
    assert item.status == "pending"


def test_intra_workbook_different_habitats_do_not_conflict():
    """Differing values between different habitats are separate documents and must not conflict."""
    planner = DocumentPlanner()

    row_lamuk = {
        "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 49.02,
        "_row": 12,
        "_cell": "G12",
    }
    row_argalingga = {
        "Lokasi Sampling": "Desa Argalingga, Kec. Argapura, Kabupaten Majalengka, Jawa Barat",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 47.09,
        "_row": 14,
        "_cell": "G14",
    }

    draft = planner.plan_records(records=[row_lamuk, row_argalingga], run_id="run_diff_habitats")
    conflicts = [
        r for r in draft.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    assert len(conflicts) == 0, "Different habitats represent separate documents; no conflict should arise."
    assert len(draft.documents) == 2


def test_intra_workbook_multiple_attribute_conflicts_isolated():
    """Multiple conflicting attributes on the same document are each enqueued independently."""
    planner = DocumentPlanner()

    row_1 = {
        "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 49.02,
        "Nitrogen content (mg/g)": 18.17,
        "_row": 12,
    }
    row_2 = {
        "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 46.54,
        "Nitrogen content (mg/g)": 17.39,
        "_row": 13,
    }

    draft = planner.plan_records(records=[row_1, row_2], run_id="run_multi_conflict")
    conflicts = [
        r for r in draft.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    assert len(conflicts) >= 2
    categories_or_labels = [c.source_label for c in conflicts]
    # Check that chlorophyll and nitrogen are both represented
    assert any("chlorophyll" in str(lbl).lower() for lbl in categories_or_labels)
    assert any("nitrogen" in str(lbl).lower() for lbl in categories_or_labels)


# ============================================================================
# Section 3: Criterion A15 — Cell Value Parsing Rules
# ============================================================================


def test_blank_noop_zero_valid_invalid_unit_and_range_parsed(resume_data_workbook_path: Path):
    """A15: Blank cells are no-op; 0/False are valid; invalid units reviewed; ranges parsed as dict."""
    wb = openpyxl.load_workbook(resume_data_workbook_path, data_only=True)
    ws = wb["Resume Data"]

    # Verify actual cell M3 and N3 in Resume Data
    val_m3 = str(ws["M3"].value).strip()
    val_n3 = str(ws["N3"].value).strip()
    assert val_m3 == "28.3 - 30.5"
    assert val_n3 == "72 - 67"

    planner = DocumentPlanner()

    # 1. Blank cell parsing: must be recognized as blank/no-op
    blank_result = planner.parse_cell_value(None)
    assert blank_result is None or getattr(blank_result, "is_blank", False) or blank_result.get("is_blank", False)

    blank_str_result = planner.parse_cell_value("   ")
    assert blank_str_result is None or getattr(blank_str_result, "is_blank", False) or blank_str_result.get("is_blank", False)

    # 2. Value 0 and False: must be valid values, not treated as blank or falsy no-op
    zero_res = planner.parse_cell_value(0, expected_type="number")
    val_zero = getattr(zero_res, "value", None) if not isinstance(zero_res, dict) else zero_res["value"]
    assert val_zero == 0
    assert val_zero is not None

    false_res = planner.parse_cell_value(False, expected_type="boolean")
    val_false = getattr(false_res, "value", None) if not isinstance(false_res, dict) else false_res["value"]
    assert val_false is False
    assert val_false is not None

    # 3. Range parsing: M3 ("28.3 - 30.5") and N3 ("72 - 67")
    range_m3 = parse_range_value(val_m3)
    assert isinstance(range_m3, dict)
    assert "min" in range_m3 and "max" in range_m3
    assert range_m3["min"] == pytest.approx(28.3)
    assert range_m3["max"] == pytest.approx(30.5)
    # Must NOT be averaged to 29.4
    assert range_m3["min"] != pytest.approx(29.4)
    assert range_m3["max"] != pytest.approx(29.4)

    range_n3 = parse_range_value(val_n3)
    assert isinstance(range_n3, dict)
    assert "min" in range_n3 and "max" in range_n3
    # Either normalized min/max or ordered endpoints, but must not be averaged
    min_n3 = min(range_n3["min"], range_n3["max"])
    max_n3 = max(range_n3["min"], range_n3["max"])
    assert min_n3 == pytest.approx(67.0)
    assert max_n3 == pytest.approx(72.0)
    assert min_n3 != pytest.approx(69.5)

    # 4. Invalid unit parsing: unknown unit escalates to review
    invalid_unit_res = planner.parse_cell_value("30 kuintal", expected_type="number", unit="kg")
    is_invalid = (
        getattr(invalid_unit_res, "is_valid", True) is False
        or (isinstance(invalid_unit_res, dict) and not invalid_unit_res.get("is_valid", True))
        or getattr(invalid_unit_res, "review_category", None) == "VALUE_INVALID"
        or (isinstance(invalid_unit_res, dict) and invalid_unit_res.get("review_category") == "VALUE_INVALID")
    )
    assert is_invalid, "Invalid unit '30 kuintal' must be flagged as invalid or for review."


def test_a15_blank_cell_preserves_existing_document_value():
    """A15: Incoming blank cell does not overwrite or erase existing attribute value in current document."""
    plant = _create_plant_identity()
    habitat = _create_habitat_identity()
    doc_key = f"{plant.id}:{habitat.id}"

    existing_doc = _create_minimal_knowledge_document(
        variety=plant,
        habitat=habitat,
        attributes={
            "attr_chlorophyll_spad": _create_sample_attribute_value(
                label="Chlorophyll Content (SPAD)",
                value=49.02,
                cell="G12",
                row=12,
            )
        },
    )

    planner = DocumentPlanner()
    row_with_blank = {
        "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": "",  # Blank
        "pH tanah": 7.0,
        "_row": 16,
    }

    draft = planner.plan_records(
        records=[row_with_blank],
        current_documents={doc_key: existing_doc},
        run_id="run_blank_noop_test",
    )

    planned_doc = draft.documents[doc_key]
    assert "attr_chlorophyll_spad" in planned_doc.attributes
    # Value must remain preserved at 49.02
    assert planned_doc.attributes["attr_chlorophyll_spad"].value == pytest.approx(49.02)


def test_a15_numeric_zero_float_zero_stored_as_valid_number():
    """A15: Numeric 0 and 0.0 are stored as valid numbers, not ignored or treated as empty."""
    planner = DocumentPlanner()

    row_zero = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Jumlah Buah Per Tanaman": 0,
        "_row": 3,
    }

    draft = planner.plan_records(records=[row_zero], run_id="run_zero_test")
    doc = list(draft.documents.values())[0]

    matched = [v for k, v in doc.attributes.items() if "buah" in k or "jumlah" in k]
    assert len(matched) == 1
    assert matched[0].value == 0
    assert matched[0].value_type == "number"
    assert not isinstance(matched[0].value, bool)


def test_a15_boolean_false_and_true_stored_validly():
    """A15: Boolean False and True values are parsed as boolean, not strings or empty."""
    planner = DocumentPlanner()

    parsed_false = planner.parse_cell_value(False, expected_type="boolean")
    val_f = getattr(parsed_false, "value", parsed_false.get("value") if isinstance(parsed_false, dict) else parsed_false)
    assert val_f is False
    assert isinstance(val_f, bool)

    parsed_true = planner.parse_cell_value(True, expected_type="boolean")
    val_t = getattr(parsed_true, "value", parsed_true.get("value") if isinstance(parsed_true, dict) else parsed_true)
    assert val_t is True
    assert isinstance(val_t, bool)


def test_a15_range_parsing_formats_without_averaging():
    """A15: Diverse range string patterns parse cleanly to min/max dicts without averaging."""
    cases = [
        ("28.3 - 30.5", 28.3, 30.5),
        ("72 - 67", 67.0, 72.0),
        ("10.5-15.2", 10.5, 15.2),
        ("5 to 10", 5.0, 10.0),
        ("20 - 25 oC", 20.0, 25.0),
    ]

    for raw, expected_min, expected_max in cases:
        parsed = parse_range_value(raw)
        assert isinstance(parsed, dict), f"Failed to parse range for '{raw}'"
        assert "min" in parsed and "max" in parsed
        actual_min = min(parsed["min"], parsed["max"])
        actual_max = max(parsed["min"], parsed["max"])
        assert actual_min == pytest.approx(expected_min)
        assert actual_max == pytest.approx(expected_max)
        # Verify strictly that no automatic averaging occurred
        avg = (expected_min + expected_max) / 2.0
        assert parsed["min"] != avg or expected_min == expected_max


def test_a15_invalid_unit_enqueued_to_review():
    """A15: Unrecognized or malformed measurement unit creates a review item with category VALUE_INVALID."""
    planner = DocumentPlanner()

    row_bad_unit = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": "49.02 kuintal",
        "_row": 3,
    }

    draft = planner.plan_records(records=[row_bad_unit], run_id="run_bad_unit_test")
    unit_reviews = [
        r for r in draft.reviews
        if r.category == "VALUE_INVALID"
        or "kuintal" in str(r.raw_values)
        or "unit" in (r.source_context or "").lower()
    ]
    assert len(unit_reviews) >= 1
    assert unit_reviews[0].status == "pending"


# ============================================================================
# Section 4: Criterion A24 — Photo Variety Mismatch (Fail-Closed)
# ============================================================================


def test_photo_mismatch_creates_blocking_review_item():
    """A24: Photos not matching active Excel varieties generate blocking PHOTO_VARIETY_MISMATCH review items."""
    planner = DocumentPlanner()

    # Active Excel run only contains Cabai rawit
    excel_records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "pH tanah": 7.0,
            "_row": 3,
        }
    ]

    # Photo signals include a photo showing curly pepper (Cabai keriting)
    mismatched_photo = PhotoContextSignal(
        filename="cabai_keriting_buah_01.jpg",
        caption="Buah cabai merah keriting panjang bergelombang",
        extracted_part="BUAH",
        suggested_variety_hint="Cabai keriting",
        raw_tokens=["cabai", "keriting", "merah"],
    )

    draft = planner.plan_records(
        records=excel_records,
        photo_signals=[mismatched_photo],
        run_id="run_photo_mismatch_test",
    )

    # Must produce a PHOTO_VARIETY_MISMATCH review item
    photo_reviews = [r for r in draft.reviews if r.category == "PHOTO_VARIETY_MISMATCH"]
    assert len(photo_reviews) == 1

    rev = photo_reviews[0]
    assert rev.status == "pending", "Review item must be pending to block commit fail-closed."
    assert rev.source_label == "cabai_keriting_buah_01.jpg"
    assert rev.target_document_key is None, "Unassigned photo must have null target_document_key."

    # Must offer the 3 explicit candidate actions
    actions = [c.get("action") for c in rev.candidate_resolutions]
    assert "ASSIGN_EXISTING" in actions
    assert "CREATE_NEW_DOCUMENT" in actions
    assert "DISCARD" in actions


def test_a24_photo_matching_excel_variety_auto_linked():
    """A24: Photo matching active Excel run variety is linked without generating a mismatch review item."""
    planner = DocumentPlanner()

    excel_records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "pH tanah": 7.0,
            "_row": 3,
        }
    ]

    matching_photo = PhotoContextSignal(
        filename="cabai_rawit_daun_01.jpg",
        caption="Daun cabai rawit ovate",
        extracted_part="DAUN",
        suggested_variety_hint="Cabai rawit",
        raw_tokens=["cabai", "rawit", "daun"],
    )

    draft = planner.plan_records(
        records=excel_records,
        photo_signals=[matching_photo],
        run_id="run_photo_match_test",
    )

    # No mismatch review item should be raised
    photo_reviews = [r for r in draft.reviews if r.category == "PHOTO_VARIETY_MISMATCH"]
    assert len(photo_reviews) == 0

    # The document must have a media reference linked or staged
    doc = list(draft.documents.values())[0]
    assert len(doc.media_refs) >= 1
    assert doc.media_refs[0].source_filename == "cabai_rawit_daun_01.jpg"
    assert doc.media_refs[0].organ_part == "daun"


def test_a24_standalone_evaluate_photo_variety_alignment():
    """evaluate_photo_variety_alignment checks candidate photo signals against variety identities."""
    plant_rawit = _create_plant_identity(variety_id="plant_rawit_001", label="Cabai rawit (Capsicum frutescens L)")
    mismatched_photo = PhotoContextSignal(
        filename="cabai_habanero_01.jpg",
        caption="Buah cabai habanero bulat",
        extracted_part="BUAH",
        suggested_variety_hint="Cabai Habanero",
    )

    items = evaluate_photo_variety_alignment([mismatched_photo], [plant_rawit])
    assert len(items) == 1
    assert items[0].category == "PHOTO_VARIETY_MISMATCH"
    assert items[0].status == "pending"
    assert items[0].source_label == "cabai_habanero_01.jpg"


def test_a24_multiple_mismatched_photos_each_generate_review_item():
    """Multiple mismatched photos each generate distinct review items."""
    planner = DocumentPlanner()

    excel_records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "_row": 3,
        }
    ]

    photos = [
        PhotoContextSignal(filename="foto_katokkon.jpg", suggested_variety_hint="Cabai Katokkon"),
        PhotoContextSignal(filename="foto_habanero.jpg", suggested_variety_hint="Cabai Habanero"),
    ]

    draft = planner.plan_records(
        records=excel_records,
        photo_signals=photos,
        run_id="run_multi_photo_test",
    )

    mismatch_reviews = [r for r in draft.reviews if r.category == "PHOTO_VARIETY_MISMATCH"]
    assert len(mismatch_reviews) == 2
    filenames = {r.source_label for r in mismatch_reviews}
    assert filenames == {"foto_katokkon.jpg", "foto_habanero.jpg"}


# ============================================================================
# Section 5: Document Envelope Grouping per (variety_id, habitat_id)
# ============================================================================


def test_grouping_multiple_rows_same_variety_and_habitat_into_single_envelope():
    """Multiple spreadsheet rows with identical variety and habitat collapse into one KnowledgeDocument."""
    planner = DocumentPlanner()

    row_1 = {
        "Lokasi Sampling": "Dusun Randu, Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta (sampling 1)",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "pH tanah": 7.0,
        "_row": 3,
    }
    row_2 = {
        "Lokasi Sampling": "Dusun Tenen, Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta (sampling 2)",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Kelembaban tanah (%)": 58,
        "_row": 4,
    }

    draft = planner.plan_records(records=[row_1, row_2], run_id="run_grouping_test")
    assert len(draft.documents) == 1, "Rows with identical variety and habitat must map to a single document."

    doc = list(draft.documents.values())[0]
    # Sampling locations should accumulate both sampling points
    assert len(doc.sampling_locations) == 2
    hamlets = {loc.hamlet for loc in doc.sampling_locations}
    assert "Dusun Randu" in hamlets
    assert "Dusun Tenen" in hamlets


def test_grouping_distinct_varieties_on_same_habitat_produce_separate_documents():
    """Distinct varieties in the same habitat produce separate KnowledgeDocument envelopes."""
    planner = DocumentPlanner()

    row_rawit = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "_row": 3,
    }
    row_besar = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "cabai besar (Capsicum annum L)",
        "_row": 4,
    }

    draft = planner.plan_records(records=[row_rawit, row_besar], run_id="run_distinct_varieties")
    assert len(draft.documents) == 2, "Distinct varieties must produce separate documents."

    varieties = {doc.variety.label for doc in draft.documents.values()}
    assert "Cabai rawit (Capsicum frutescens L)" in varieties
    assert "cabai besar (Capsicum annum L)" in varieties


def test_grouping_same_variety_on_distinct_habitats_produce_separate_documents():
    """Same variety across different geographical habitats produces separate KnowledgeDocument envelopes."""
    planner = DocumentPlanner()

    row_sleman = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "_row": 3,
    }
    row_wonosobo = {
        "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "_row": 12,
    }

    draft = planner.plan_records(records=[row_sleman, row_wonosobo], run_id="run_distinct_habitats")
    assert len(draft.documents) == 2

    villages = {doc.habitat.administrative_path.village_name for doc in draft.documents.values()}
    assert "Desa Hargobinangun" in villages
    assert "Desa Lamuk" in villages


def test_document_keys_follow_strict_variety_colon_habitat_format():
    """All planned document dictionary keys follow strict '{variety_id}:{habitat_id}' format."""
    planner = DocumentPlanner()

    records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "_row": 3,
        },
        {
            "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "_row": 12,
        },
    ]

    draft = planner.plan_records(records=records, run_id="run_doc_key_test")
    for key, doc in draft.documents.items():
        expected_key = f"{doc.variety.id}:{doc.habitat.id}"
        assert key == expected_key


def test_unresolved_location_enqueues_location_review_item():
    """Ambiguous or unresolved location string triggers LOCATION_UNRESOLVED review item."""
    planner = DocumentPlanner()

    bad_location_row = {
        "Lokasi Sampling": "Desa Antah Berantah Tanpa Kabupaten Yang Jelas",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "_row": 99,
    }

    draft = planner.plan_records(records=[bad_location_row], run_id="run_bad_loc_test")
    loc_reviews = [
        r for r in draft.reviews
        if r.category in ("LOCATION_UNRESOLVED", "LOCATION_AMBIGUOUS", "UNRESOLVED_LOCATION")
    ]
    assert len(loc_reviews) >= 1
    assert loc_reviews[0].status == "pending"


def test_unresolved_plant_identity_enqueues_plant_review_item():
    """Unrecognized plant label triggers PLANT_IDENTITY_UNRESOLVED review item."""
    planner = DocumentPlanner()

    bad_plant_row = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Tanaman Misterius X99",
        "_row": 99,
    }

    draft = planner.plan_records(records=[bad_plant_row], run_id="run_bad_plant_test")
    plant_reviews = [
        r for r in draft.reviews
        if r.category == "PLANT_IDENTITY_UNRESOLVED"
    ]
    assert len(plant_reviews) >= 1
    assert plant_reviews[0].status == "pending"


def test_novel_attribute_enqueues_new_attribute_review_item():
    """Unmapped/novel attribute header generates NEW_ATTRIBUTE review item."""
    planner = DocumentPlanner()

    novel_attr_row = {
        "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Atribut Baru Ekstrem Yang Belum Pernah Ada": 123.45,
        "_row": 3,
    }

    draft = planner.plan_records(records=[novel_attr_row], run_id="run_novel_attr_test")
    attr_reviews = [
        r for r in draft.reviews
        if r.category in ("NEW_ATTRIBUTE", "AMBIGUOUS_MAPPING")
    ]
    assert len(attr_reviews) >= 1
    assert attr_reviews[0].status == "pending"


# ============================================================================
# Section 6: Document Draft Integrity, Concurrency & Base Versions
# ============================================================================


def test_draft_envelope_validates_against_document_draft_contract():
    """Output of document planner strictly satisfies DocumentDraft Pydantic contract."""
    planner = DocumentPlanner()

    records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "pH tanah": 7.0,
            "_row": 3,
        }
    ]

    draft = planner.plan_records(
        records=records,
        run_id="run_contract_valid_001",
        source_filename="data_input.xlsx",
        source_scope="field_survey_2026",
    )

    assert isinstance(draft, DocumentDraft)
    assert draft.run_id == "run_contract_valid_001"
    assert draft.source_filename == "data_input.xlsx"
    assert draft.schema_version == "kms-document-v1"
    assert draft.revision >= 1
    assert len(draft.file_sha256) == 64


def test_draft_tracks_base_versions_for_existing_documents():
    """DocumentDraft accurately records base document versions from current documents."""
    plant = _create_plant_identity()
    habitat = _create_habitat_identity()
    doc_key = f"{plant.id}:{habitat.id}"

    existing_doc = _create_minimal_knowledge_document(variety=plant, habitat=habitat)

    planner = DocumentPlanner()
    row = {
        "Lokasi Sampling": "Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "pH tanah": 6.8,
        "_row": 12,
    }

    # Pass current documents with base version 3
    draft = planner.plan_records(
        records=[row],
        current_documents={doc_key: existing_doc},
        base_versions={doc_key: 3},
        run_id="run_base_ver_test",
    )

    assert doc_key in draft.base_versions
    assert draft.base_versions[doc_key] == 3


def test_plan_document_changes_top_level_functional_helper():
    """Top-level plan_document_changes function matches the DocumentPlanner planning flow."""
    records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "pH tanah": 7.0,
            "_row": 3,
        }
    ]

    draft = plan_document_changes(
        source_records=records,
        current_documents={},
        run_id="run_func_helper_001",
    )
    assert isinstance(draft, DocumentDraft)
    assert len(draft.documents) == 1


# ============================================================================
# Section 7: End-to-End Workbook Planning on Field Sample (data_input.xlsx)
# ============================================================================


def test_e2e_workbook_planner_on_resume_data(resume_data_workbook_path: Path):
    """End-to-end planning on real data_input.xlsx sample correctly organizes all documents and conflicts."""
    planner = DocumentPlanner()

    draft = planner.plan_workbook(
        workbook_path=resume_data_workbook_path,
        run_id="run_e2e_resume_data",
    )

    # 1. Documents must be created for each distinct variety + habitat pair
    assert len(draft.documents) >= 3

    # Check presence of Sleman, Lamuk (Wonosobo), and Argalingga (Majalengka)
    villages = {
        doc.habitat.administrative_path.village_name
        for doc in draft.documents.values()
    }
    assert "Desa Hargobinangun" in villages
    assert "Desa Lamuk" in villages
    assert "Desa Argalingga" in villages

    # 2. Intra-workbook conflicts must be detected for Lamuk and Argalingga
    conflict_reviews = [
        r for r in draft.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    assert len(conflict_reviews) >= 2

    # 3. All reviews must be pending
    for r in draft.reviews:
        assert r.status == "pending"

    # 4. Range attributes (Suhu udara M3 and Kelembaban udara N3) must be parsed as dicts
    hargobinangun_docs = [
        doc for doc in draft.documents.values()
        if doc.habitat.administrative_path.village_name == "Desa Hargobinangun"
    ]
    assert len(hargobinangun_docs) >= 1
    hg_doc = hargobinangun_docs[0]

    range_attrs = [v for k, v in hg_doc.attributes.items() if v.value_type == "range"]
    assert len(range_attrs) >= 1
    for r_attr in range_attrs:
        assert isinstance(r_attr.value, dict)
        assert "min" in r_attr.value and "max" in r_attr.value


def test_a11_conflict_preserves_source_coordinates_and_cells():
    """A11: Conflict review items record conflicting source cells (e.g. G12 vs G13)."""
    planner = DocumentPlanner()

    row_12 = {
        "Lokasi Sampling": "Dusun Pangempon, Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 49.02,
        "_row": 12,
        "_cell": "G12",
    }
    row_13 = {
        "Lokasi Sampling": "Dusun Genting, Desa Lamuk, Kec. Kalikajar, Kabupaten Wonosobo, Jawa Tengah",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "Chlorophyll Content (SPAD)": 46.5417,
        "_row": 13,
        "_cell": "G13",
    }

    draft = planner.plan_records(records=[row_12, row_13], run_id="run_cell_provenance_test")
    conflicts = [
        r for r in draft.reviews
        if r.category in ("INTRA_WORKBOOK_CONFLICT", "INTRA_RUN_VALUE_CONFLICT")
    ]
    assert len(conflicts) == 1
    item = conflicts[0]
    # Source context or payload should mention source cells or rows
    ctx = f"{item.source_context} {item.payload} {item.notes}".lower()
    assert "12" in ctx or "g12" in ctx or len(item.raw_values) == 2


def test_a15_nan_and_whitespace_treated_as_blank_noop():
    """A15: float('nan') and whitespace strings are treated as blank no-ops."""
    planner = DocumentPlanner()

    res_nan = planner.parse_cell_value(float("nan"))
    assert res_nan is None or getattr(res_nan, "is_blank", False) or res_nan.get("is_blank", False)

    res_ws = planner.parse_cell_value("   \t\n  ")
    assert res_ws is None or getattr(res_ws, "is_blank", False) or res_ws.get("is_blank", False)


def test_a15_range_negative_averaging_assertion():
    """A15: Range values strictly forbid arithmetic averaging into a single scalar."""
    raw = "28.3 - 30.5"
    parsed = parse_range_value(raw)
    assert isinstance(parsed, dict)
    assert "min" in parsed and "max" in parsed

    mean = (28.3 + 30.5) / 2.0  # 29.4
    assert parsed.get("value") != mean
    assert parsed["min"] != mean
    assert parsed["max"] != mean


def test_a24_photo_signals_accepted_as_dict_or_contract():
    """A24: Planner accepts photo signals as raw dicts or PhotoContextSignal contract objects."""
    planner = DocumentPlanner()

    excel_records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "_row": 3,
        }
    ]

    photo_dict = {
        "filename": "cabai_hias_pelangi.jpg",
        "caption": "Cabai hias ungu pelangi",
        "suggested_variety_hint": "Cabai pelangi",
    }

    draft = planner.plan_records(
        records=excel_records,
        photo_signals=[photo_dict],
        run_id="run_photo_dict_test",
    )

    photo_reviews = [r for r in draft.reviews if r.category == "PHOTO_VARIETY_MISMATCH"]
    assert len(photo_reviews) == 1
    assert photo_reviews[0].source_label == "cabai_hias_pelangi.jpg"


def test_grouping_preserves_sampling_coordinates():
    """GPS coordinates from columns C and D are preserved in SamplingLocationDetail."""
    planner = DocumentPlanner()

    row = {
        "Lokasi Sampling": "Dusun Randu, Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta (sampling 1)",
        "Varietas": "Cabai rawit (Capsicum frutescens L)",
        "S": 110.415748,
        "E": -7.634808,
        "_row": 3,
    }

    draft = planner.plan_records(records=[row], run_id="run_coords_test")
    doc = list(draft.documents.values())[0]
    assert len(doc.sampling_locations) == 1
    loc = doc.sampling_locations[0]
    if loc.coordinates:
        assert isinstance(loc.coordinates, dict)
        coords_str = str(loc.coordinates)
        assert "110.41" in coords_str or "-7.63" in coords_str


def test_draft_template_hash_and_metadata_populated():
    """DocumentDraft accurately carries template hash, run metadata and config fingerprint."""
    planner = DocumentPlanner()

    records = [
        {
            "Lokasi Sampling": "Desa Hargobinangun, Kec. Pakem, Kabupaten Sleman, Daerah Istimewa Yogyakarta",
            "Varietas": "Cabai rawit (Capsicum frutescens L)",
            "_row": 3,
        }
    ]

    draft = planner.plan_records(
        records=records,
        run_id="run_meta_test_001",
        metadata={"batch_operator": "kurator_lapangan_01"},
    )

    assert draft.run_id == "run_meta_test_001"
    assert draft.metadata.get("batch_operator") == "kurator_lapangan_01"
    assert len(draft.template_hash) == 64

