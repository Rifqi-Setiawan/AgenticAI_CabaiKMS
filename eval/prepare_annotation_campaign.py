"""Build a fail-closed, frozen human-annotation campaign from one manifest."""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from eval.create_mapping_annotations import ANNOTATION_COLUMNS, HUMAN_COLUMNS, create_annotation_table
from src.ingestion.runtime_source import prepare_legacy_runtime_source
from src.schema.canonical import CanonicalSchema
from src.schema.evaluation_config import EvaluationRunConfig
from src.schema.gold_mapping import MappingIdentityKind, build_mapping_item_identities
from src.ui.pipeline_runner import run_pipeline_ui
from src.ui import pipeline_runner

CAMPAIGN_VERSION = "schema-mapping-validation-campaign-v1"
GUIDELINE_VERSION = "annotation-guidelines-v1"
SUPPORTED_FORMATS = {"row-oriented", "transposed"}
PREDICTION_COLUMNS = {
    "proposed_target_canonical_key", "predicted_row", "mapping_method", "confidence",
    "exact_name_status", "verifier_status", "verifier_warnings", "verifier_hard_issues",
    "retrieval_target_rank", "retrieval_target_distance", "retrieval_top1_top2_margin",
    "retrieval_target_vs_top1_gap", "acceptance_status",
}
BLINDED_COLUMNS = [
    "mapping_item_id", "source_attribute_display", "source_attribute", "source_context",
    *HUMAN_COLUMNS, "annotation_version", "annotation_source",
    "mapping_identity_kind", "mapping_identity_value", "mapping_identity_issue",
    "source_file_name", "source_file_sha256", "source_sheet", "source_format",
    "source_attribute_id", "source_backend", "retrieval_backend", "retrieval_k",
    "schema_version", "template_hash", "mapping_verification_version",
    "embedding_model_name", "evaluation_config_fingerprint",
]
IMMUTABLE_COLUMNS = [column for column in BLINDED_COLUMNS if column not in HUMAN_COLUMNS]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _json_bytes(value: dict) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def _manifest_config(manifest: dict) -> EvaluationRunConfig:
    return EvaluationRunConfig(
        source_backend=manifest["source_backend"], retrieval_backend=manifest["retrieval_backend"],
        retrieval_k=manifest["retrieval_k"],
        canonical_schema_version=manifest["canonical_schema_version"],
        canonical_template_hash=manifest["canonical_template_hash"],
        mapping_verification_version=manifest["mapping_verification_version"],
        embedding_model_name=manifest["embedding_model_name"],
    )


def assess_campaign_readiness(
    manifest_path: Path, *, project_root: Path = PROJECT_ROOT,
    source_loader: Callable = prepare_legacy_runtime_source,
) -> dict:
    """Inspect actual inputs and return a deterministic, fail-closed report."""
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    schema = CanonicalSchema.from_template()
    blockers: list[str] = []
    expected_config = _manifest_config(manifest)
    declared_fingerprint = manifest.get("frozen_evaluation_config_fingerprint", "")
    if declared_fingerprint != expected_config.fingerprint:
        blockers.append("FROZEN_EVALUATION_CONFIG_FINGERPRINT_MISMATCH")
    if manifest.get("canonical_schema_version") != schema.schema_version:
        blockers.append("CANONICAL_SCHEMA_VERSION_MISMATCH")
    if manifest.get("canonical_template_hash") != schema.template_hash:
        blockers.append("CANONICAL_TEMPLATE_HASH_MISMATCH")

    validation_entries = [item for item in manifest.get("validation_workbooks", []) if item.get("part_of_validation") is True]
    if not validation_entries:
        blockers.append("NO_VALIDATION_WORKBOOKS")
    counts = {
        "validation_workbook_count": len(validation_entries), "sheet_count": 0,
        "total_source_attribute_count": 0, "annotatable_item_count": 0,
        "stable_identity_count": 0, "unavailable_identity_count": 0,
    }
    item_ids: list[str] = []
    seen_entry_ids: set[str] = set()
    for entry in validation_entries:
        entry_id = str(entry.get("entry_id", "")).strip()
        if not entry_id or entry_id in seen_entry_ids:
            blockers.append(f"INVALID_OR_DUPLICATE_ENTRY_ID:{entry_id or '<blank>'}")
        seen_entry_ids.add(entry_id)
        if (entry.get("source_backend"), entry.get("retrieval_backend"), entry.get("retrieval_k")) != (
            manifest.get("source_backend"), manifest.get("retrieval_backend"), manifest.get("retrieval_k")
        ):
            blockers.append(f"MIXED_EVALUATION_CONFIGURATION:{entry_id}")
        source_format = entry.get("source_format")
        if source_format not in SUPPORTED_FORMATS:
            blockers.append(f"UNSUPPORTED_SOURCE_FORMAT:{entry_id}:{source_format}")
            continue
        source = project_root / entry.get("path", "")
        if not source.is_file():
            blockers.append(f"VALIDATION_FILE_MISSING:{entry_id}:{entry.get('path', '')}")
            continue
        actual_hash = _sha256(source)
        if actual_hash != str(entry.get("sha256", "")).lower():
            blockers.append(f"WORKBOOK_HASH_MISMATCH:{entry_id}")
            continue
        try:
            available_sheets = set(pd.ExcelFile(source).sheet_names)
        except Exception as exc:
            blockers.append(f"WORKBOOK_UNREADABLE:{entry_id}:{type(exc).__name__}")
            continue
        sheets = entry.get("sheets", [])
        missing_sheets = sorted(set(sheets) - available_sheets)
        if missing_sheets:
            blockers.append(f"REQUIRED_SHEET_MISSING:{entry_id}:{'|'.join(missing_sheets)}")
            continue
        entry_counts = {"annotatable_attributes": 0, "stable_mapping_identities": 0, "unavailable_identities": 0}
        for sheet in sheets:
            counts["sheet_count"] += 1
            try:
                bundle = source_loader(source, sheet, source_format=source_format)
                counts["total_source_attribute_count"] += len(bundle.all_attributes)
                identities = build_mapping_item_identities(
                    source_file_sha256=actual_hash, source_sheet=sheet, source_format=source_format,
                    source_items=[(attribute.source_attribute_id, attribute.display_name) for attribute in bundle.schema_attributes],
                )
            except Exception as exc:
                blockers.append(f"SOURCE_PARSE_FAILED:{entry_id}:{sheet}:{type(exc).__name__}")
                continue
            stable = [item.mapping_item_id for item in identities if item.identity_kind is not MappingIdentityKind.UNAVAILABLE]
            unavailable = sum(item.identity_kind is MappingIdentityKind.UNAVAILABLE for item in identities)
            entry_counts["annotatable_attributes"] += len(identities)
            entry_counts["stable_mapping_identities"] += len(stable)
            entry_counts["unavailable_identities"] += unavailable
            counts["annotatable_item_count"] += len(identities)
            counts["stable_identity_count"] += len(stable)
            counts["unavailable_identity_count"] += unavailable
            item_ids.extend(item_id for item_id in stable if item_id)
            if unavailable:
                blockers.append(f"IDENTITY_UNAVAILABLE:{entry_id}:{sheet}:{unavailable}")
        for name, actual in entry_counts.items():
            if entry.get("counts", {}).get(name) != actual:
                blockers.append(f"DECLARED_COUNT_MISMATCH:{entry_id}:{name}")

    duplicate_count = len(item_ids) - len(set(item_ids))
    if duplicate_count:
        blockers.append(f"DUPLICATE_MAPPING_ITEM_ID:{duplicate_count}")
    return {
        "campaign_version": manifest.get("campaign_version"), **counts,
        "duplicate_mapping_item_id_count": duplicate_count,
        "frozen_evaluation_config_fingerprint": declared_fingerprint,
        "canonical_schema_version": manifest.get("canonical_schema_version"),
        "canonical_template_hash": manifest.get("canonical_template_hash"),
        "campaign_ready": not blockers, "blockers": sorted(set(blockers)),
    }


def _reference_frame(schema: CanonicalSchema) -> pd.DataFrame:
    return pd.DataFrame([{
        "canonical_key": row.canonical_key, "label": row.label, "domain": row.domain,
        "example_values": " | ".join(row.contoh_nilai), "alternate_labels": " | ".join(row.alt_labels),
    } for row in schema.rows])


def _style_workbook(path: Path, *, blinded: bool) -> None:
    workbook = load_workbook(path)
    sheet = workbook["Annotations"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    human_indexes = {headers[column] for column in HUMAN_COLUMNS if column in headers}
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            editable = blinded and cell.column in human_indexes
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if editable else "F3F4F6")
            cell.protection = Protection(locked=not editable)
    for cell in sheet[1]:
        sheet.column_dimensions[cell.column_letter].width = 18
    for name in ("source_attribute_display", "source_context", "notes"):
        if name in headers:
            sheet.column_dimensions[sheet.cell(1, headers[name]).column_letter].width = 34
    if blinded:
        status = DataValidation(type="list", formula1='"ONE_TO_ONE,NO_MATCH,AMBIGUOUS,COMPOSITE,EXCLUDE"')
        round_number = DataValidation(type="whole", operator="equal", formula1="1")
        sheet.add_data_validation(status)
        sheet.add_data_validation(round_number)
        status.add(f"{sheet.cell(2, headers['gold_status']).coordinate}:{sheet.cell(sheet.max_row, headers['gold_status']).coordinate}")
        round_number.add(f"{sheet.cell(2, headers['annotation_round']).coordinate}:{sheet.cell(sheet.max_row, headers['annotation_round']).coordinate}")
    sheet.freeze_panes = "E2" if blinded else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.protection.sheet = True
    reference = workbook["Canonical Reference"]
    for cell in reference[1]:
        cell.fill = PatternFill("solid", fgColor="5B9BD5")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for row in reference.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate([28, 34, 16, 42, 38], start=1):
        reference.column_dimensions[reference.cell(1, index).column_letter].width = width
    reference.freeze_panes = "A2"
    reference.auto_filter.ref = reference.dimensions
    reference.sheet_view.showGridLines = False
    reference.protection.sheet = True
    workbook.save(path)


def _write_workbook(path: Path, annotations: pd.DataFrame, schema: CanonicalSchema, *, blinded: bool) -> None:
    with pd.ExcelWriter(path) as writer:
        annotations.to_excel(writer, sheet_name="Annotations", index=False)
        _reference_frame(schema).to_excel(writer, sheet_name="Canonical Reference", index=False)
    _style_workbook(path, blinded=blinded)


def _git_revision(project_root: Path) -> str | None:
    try:
        result = subprocess.run(["git", "rev-parse", "HEAD"], cwd=project_root, capture_output=True, text=True, check=True)
        return result.stdout.strip() or None
    except (OSError, subprocess.CalledProcessError):
        return None


def prepare_campaign(
    manifest_path: Path, output_dir: Path, *, project_root: Path = PROJECT_ROOT,
    guideline_path: Path | None = None, pipeline_call: Callable = run_pipeline_ui,
) -> tuple[Path, Path, Path, Path, Path]:
    """Write readiness, master, A/B, and cryptographic campaign metadata."""
    output_dir.mkdir(parents=True, exist_ok=True)
    readiness_path = output_dir / "readiness.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    readiness = assess_campaign_readiness(manifest_path, project_root=project_root)
    readiness_path.write_bytes(_json_bytes(readiness))
    if not readiness["campaign_ready"]:
        raise ValueError(f"campaign is not ready: {readiness['blockers']}")
    master_path = output_dir / "master" / "master_evaluation.xlsx"
    output_a = output_dir / "round1" / "annotator_A_round1.xlsx"
    output_b = output_dir / "round1" / "annotator_B_round1.xlsx"
    metadata_path = output_dir / "metadata.json"
    collisions = [path for path in (master_path, output_a, output_b, metadata_path) if path.exists()]
    if collisions:
        raise FileExistsError(f"refusing to overwrite campaign artifact(s): {collisions}")
    for path in (master_path, output_a, output_b):
        path.parent.mkdir(parents=True, exist_ok=True)

    schema = CanonicalSchema.from_template()
    expected_config = _manifest_config(manifest)
    frames: list[pd.DataFrame] = []
    original_orchestrator_call = pipeline_runner.run_pipeline
    with tempfile.TemporaryDirectory(prefix="cabai-kms-annotation-") as temporary:
        if pipeline_call is run_pipeline_ui:
            isolated_db = Path(temporary) / "orchestrator.sqlite"
            def isolated_orchestrator(*args, **kwargs):
                return original_orchestrator_call(*args, db_path=isolated_db, **kwargs)
            pipeline_runner.run_pipeline = isolated_orchestrator
        try:
            for entry in manifest["validation_workbooks"]:
                if not entry["part_of_validation"]:
                    continue
                source = project_root / entry["path"]
                for sheet in entry["sheets"]:
                    result = pipeline_call(
                        source, source_format=entry["source_format"], sheet_name=sheet,
                        source_backend=entry["source_backend"], retrieval_backend=entry["retrieval_backend"],
                        k=entry["retrieval_k"],
                    )
                    if result.evaluation_config_fingerprint != expected_config.fingerprint:
                        raise ValueError(f"evaluation configuration mismatch for {entry['entry_id']}:{sheet}")
                    frames.append(create_annotation_table(result))
        finally:
            pipeline_runner.run_pipeline = original_orchestrator_call

    master = pd.concat(frames, ignore_index=True).reindex(columns=ANNOTATION_COLUMNS)
    if master["mapping_item_id"].duplicated().any():
        raise ValueError("validation campaign contains duplicate mapping_item_id values")
    if len(master) != readiness["annotatable_item_count"]:
        raise ValueError("master item count differs from readiness report")
    if any(master[column].fillna("").astype(str).str.strip().ne("").any() for column in HUMAN_COLUMNS):
        raise ValueError("master human gold fields must remain blank")
    blinded = master.reindex(columns=BLINDED_COLUMNS).copy(deep=True)
    if PREDICTION_COLUMNS & set(blinded.columns):
        raise ValueError("prediction columns leaked into blinded view")
    _write_workbook(master_path, master, schema, blinded=False)
    _write_workbook(output_a, blinded, schema, blinded=True)
    _write_workbook(output_b, blinded.copy(deep=True), schema, blinded=True)

    guideline = guideline_path or project_root / manifest["annotation_guideline_path"]
    metadata = {
        "campaign_version": manifest["campaign_version"], "annotation_version": manifest["annotation_version"],
        "annotation_round": manifest["annotation_round"],
        "generated_at_utc": datetime.now(timezone.utc).isoformat(), "code_revision": _git_revision(project_root),
        "campaign_manifest": {"path": manifest_path.as_posix(), "sha256": _sha256(manifest_path)},
        "readiness": {"path": readiness_path.as_posix(), "sha256": _sha256(readiness_path)},
        "validation_workbooks": [
            {"entry_id": entry["entry_id"], "path": entry["path"], "sha256": entry["sha256"]}
            for entry in manifest["validation_workbooks"] if entry["part_of_validation"]
        ],
        "master_artifact": {"path": master_path.as_posix(), "sha256": _sha256(master_path)},
        "annotator_A_template": {"path": output_a.as_posix(), "sha256": _sha256(output_a)},
        "annotator_B_template": {"path": output_b.as_posix(), "sha256": _sha256(output_b)},
        "canonical_schema_version": manifest["canonical_schema_version"],
        "canonical_template_hash": manifest["canonical_template_hash"],
        "frozen_evaluation_config_fingerprint": expected_config.fingerprint,
        "annotation_guideline": {"path": guideline.as_posix(), "version": manifest["annotation_guideline_version"], "sha256": _sha256(guideline)},
        "prediction_blinded": True,
    }
    metadata_path.write_bytes(_json_bytes(metadata))
    return readiness_path, master_path, output_a, output_b, metadata_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, default=PROJECT_ROOT)
    parser.add_argument("--guideline", type=Path)
    args = parser.parse_args()
    for path in prepare_campaign(args.manifest, args.output_dir, project_root=args.project_root, guideline_path=args.guideline):
        print(path)


if __name__ == "__main__":
    main()
