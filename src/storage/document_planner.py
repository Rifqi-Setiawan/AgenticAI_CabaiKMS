"""Document draft planner, intra-workbook conflict detection, and review queue staging for CabaiKMS JSONB storage."""

from __future__ import annotations

import copy
from dataclasses import dataclass
import hashlib
import math
from pathlib import Path
import re
from typing import Any, Literal
import uuid

import openpyxl

from src.storage.document_contracts import (
    AttributeValue,
    CurrentSourceReference,
    DocumentDraft,
    DocumentReviewItem,
    HabitatIdentity,
    KnowledgeDocument,
    MediaReference,
    PhotoContextSignal,
    PlantIdentity,
    SamplingLocationDetail,
)
from src.storage.attribute_policy import (
    AttributeCatalog,
    AttributeDefinition,
    AttributePolicy,
    resolve_attribute,
)
from src.storage.identity_resolution import (
    AdministrativeRegistry,
    PlantCatalog,
    build_document_identity,
    resolve_plant_identity,
    resolve_sampling_location,
)

__all__ = [
    "DocumentPlanner",
    "ParsedCellValue",
    "detect_intra_workbook_conflicts",
    "evaluate_photo_variety_alignment",
    "parse_cell_value",
    "parse_range_value",
    "plan_document_changes",
    "plan_document_draft",
]


@dataclass
class ParsedCellValue:
    """Parsed representation of a spreadsheet cell value."""

    value: Any = None
    value_type: str = "text"
    is_blank: bool = False
    is_valid: bool = True
    review_category: str | None = None
    unit: str | None = None
    error: str | None = None

    def get(self, key: str, default: Any = None) -> Any:
        return getattr(self, key, default)


def parse_range_value(raw: Any) -> dict[str, float] | None:
    """Parse range string into min and max boundaries without arithmetic averaging."""
    if not isinstance(raw, str):
        return None
    raw_str = raw.strip()
    if not any(sep in raw_str.lower() for sep in ("-", "to", "sampai")):
        return None
    m = re.search(
        r"(\d+(?:\.\d*)?)\s*(?:-|to|sampai)\s*(\d+(?:\.\d*)?)",
        raw_str,
        re.IGNORECASE,
    )
    if m:
        try:
            val1 = float(m.group(1))
            val2 = float(m.group(2))
            return {"min": min(val1, val2), "max": max(val1, val2)}
        except ValueError:
            return None
    return None


def parse_cell_value(
    value: Any,
    expected_type: str | None = None,
    unit: str | None = None,
) -> ParsedCellValue:
    """Parse cell value respecting blank no-ops, booleans, ranges, numbers, and units."""
    if value is None:
        return ParsedCellValue(value=None, is_blank=True, is_valid=True)
    if isinstance(value, float) and math.isnan(value):
        return ParsedCellValue(value=None, is_blank=True, is_valid=True)
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return ParsedCellValue(value=None, is_blank=True, is_valid=True)

    if isinstance(value, bool):
        return ParsedCellValue(
            value=value, value_type="boolean", is_blank=False, is_valid=True
        )

    if isinstance(value, (int, float)):
        return ParsedCellValue(
            value=value,
            value_type="number",
            is_blank=False,
            is_valid=True,
            unit=unit,
        )

    if isinstance(value, str):
        stripped = value.strip()
        if expected_type == "boolean":
            if stripped.lower() in ("true", "ya", "yes", "1"):
                return ParsedCellValue(
                    value=True, value_type="boolean", is_blank=False, is_valid=True
                )
            if stripped.lower() in ("false", "tidak", "no", "0"):
                return ParsedCellValue(
                    value=False, value_type="boolean", is_blank=False, is_valid=True
                )

        range_dict = parse_range_value(stripped)
        if range_dict is not None:
            return ParsedCellValue(
                value=range_dict,
                value_type="range",
                is_blank=False,
                is_valid=True,
                unit=unit,
            )

        tokens = stripped.split()
        if len(tokens) >= 2:
            try:
                num_val = float(tokens[0])
                found_unit = " ".join(tokens[1:]).strip()
                norm_expected = unit.strip().lower() if unit else None
                if norm_expected is None or found_unit.lower() != norm_expected:
                    return ParsedCellValue(
                        value=num_val,
                        value_type=expected_type or "number",
                        is_blank=False,
                        is_valid=False,
                        review_category="VALUE_INVALID",
                        unit=found_unit,
                        error=f"Unrecognized or mismatched measurement unit '{found_unit}'",
                    )
                return ParsedCellValue(
                    value=num_val,
                    value_type="number",
                    is_blank=False,
                    is_valid=True,
                    unit=unit,
                )
            except ValueError:
                pass

        if expected_type == "number" or re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)", stripped):
            try:
                if "." in stripped or "e" in stripped.lower():
                    return ParsedCellValue(
                        value=float(stripped),
                        value_type="number",
                        is_blank=False,
                        is_valid=True,
                        unit=unit,
                    )
                return ParsedCellValue(
                    value=int(stripped),
                    value_type="number",
                    is_blank=False,
                    is_valid=True,
                    unit=unit,
                )
            except ValueError:
                pass

        return ParsedCellValue(
            value=stripped,
            value_type="text" if expected_type == "number" else (expected_type or "text"),
            is_blank=False,
            is_valid=True,
            unit=None if expected_type == "number" else unit,
        )

    return ParsedCellValue(
        value=value,
        value_type=expected_type or "text",
        is_blank=False,
        is_valid=True,
        unit=unit,
    )


_KNOWN_FIELD_DEFINITIONS: list[dict[str, Any]] = [
    {
        "attribute_id": "attr_chlorophyll_spad",
        "canonical_key": "chlorophyll_spad",
        "label": "Chlorophyll Content (SPAD)",
        "organ": "daun",
        "value_type": "number",
        "unit": "SPAD",
        "aliases": ["chlorophyll", "klorofil", "spad", "chlorophyll content"],
    },
    {
        "attribute_id": "attr_nitrogen_content",
        "canonical_key": "nitrogen_content",
        "label": "Nitrogen content (mg/g)",
        "organ": "daun",
        "value_type": "number",
        "unit": "mg/g",
        "aliases": ["nitrogen", "n_content"],
    },
    {
        "attribute_id": "attr_leaf_humidity",
        "canonical_key": "leaf_humidity",
        "label": "Humidity (%)",
        "organ": "daun",
        "value_type": "number",
        "unit": "%",
        "aliases": ["humidity daun", "kelembaban daun"],
    },
    {
        "attribute_id": "attr_leaf_surface_temp",
        "canonical_key": "leaf_surface_temp",
        "label": "Suhu permukaan daun(oC)",
        "organ": "daun",
        "value_type": "number",
        "unit": "oC",
        "aliases": ["suhu daun", "suhu permukaan daun"],
    },
    {
        "attribute_id": "attr_soil_humidity",
        "canonical_key": "soil_humidity",
        "label": "Kelembaban tanah (%)",
        "organ": "ekologi_tanah",
        "value_type": "number",
        "unit": "%",
        "aliases": ["kelembaban tanah"],
    },
    {
        "attribute_id": "attr_soil_ph",
        "canonical_key": "soil_ph",
        "label": "pH tanah",
        "organ": "ekologi_tanah",
        "value_type": "number",
        "unit": None,
        "aliases": ["ph tanah", "soil ph", "ph"],
    },
    {
        "attribute_id": "attr_air_temp",
        "canonical_key": "air_temp",
        "label": "Suhu udara (0C)",
        "organ": "microclimate",
        "value_type": "range",
        "unit": "0C",
        "aliases": ["suhu udara", "air temperature"],
    },
    {
        "attribute_id": "attr_air_humidity",
        "canonical_key": "air_humidity",
        "label": "Kelembaban udara (%)",
        "organ": "microclimate",
        "value_type": "range",
        "unit": "%",
        "aliases": ["kelembaban udara", "air humidity"],
    },
    {
        "attribute_id": "attr_jumlah_buah_per_tanaman",
        "canonical_key": "jumlah_buah_per_tanaman",
        "label": "Jumlah Buah Per Tanaman",
        "organ": "buah",
        "value_type": "number",
        "unit": None,
        "aliases": [
            "jumlah buah",
            "jumlah buah per tanaman",
            "buah per tanaman",
            "fruit per plant",
        ],
    },
]


def _are_values_equal(val1: Any, val2: Any) -> bool:
    """Compare two attribute values for equality."""
    if isinstance(val1, (int, float)) and isinstance(val2, (int, float)):
        if isinstance(val1, bool) or isinstance(val2, bool):
            return val1 is val2
        return math.isclose(float(val1), float(val2), rel_tol=1e-5, abs_tol=1e-5)
    if isinstance(val1, dict) and isinstance(val2, dict):
        if "min" in val1 and "max" in val1 and "min" in val2 and "max" in val2:
            return (
                math.isclose(float(val1["min"]), float(val2["min"]), rel_tol=1e-5, abs_tol=1e-5)
                and math.isclose(float(val1["max"]), float(val2["max"]), rel_tol=1e-5, abs_tol=1e-5)
            )
        return val1 == val2
    if isinstance(val1, str) and isinstance(val2, str):
        return val1.strip().lower() == val2.strip().lower()
    return val1 == val2


def detect_intra_workbook_conflicts(
    target_key: str,
    grouped_rows: list[dict[str, Any]],
    run_id: str | None = None,
) -> list[DocumentReviewItem]:
    """Detect conflicting values for the same attribute across rows for the same document."""
    if len(grouped_rows) <= 1:
        return []

    non_attribute_keys = {
        "lokasi sampling", "lokasi_sampling", "alamat", "location", "address",
        "varietas", "jenis cabai", "jenis_cabai", "plant", "variety",
        "s", "e", "coordinates", "latitude", "longitude",
        "ketinggian lahan sampling (mdpl = meter diatas permukaan laut)",
        "elevation", "elevation_mdpl",
    }

    attribute_keys: set[str] = set()
    for row in grouped_rows:
        for k in row.keys():
            if k.startswith("_") or k.strip().lower() in non_attribute_keys:
                continue
            attribute_keys.add(k)

    conflicts: list[DocumentReviewItem] = []

    for attr in sorted(attribute_keys):
        collected_values: list[Any] = []
        collected_cells: list[str] = []
        collected_rows: list[int] = []

        for row in grouped_rows:
            if attr not in row:
                continue
            raw_v = row[attr]
            parsed = parse_cell_value(raw_v)
            if parsed.is_blank:
                continue

            collected_values.append(parsed.value)
            cell_ref = (
                row.get(f"_cell_{attr}")
                or row.get("_cell")
                or (f"row_{row.get('_row')}" if row.get("_row") else "")
            )
            collected_cells.append(str(cell_ref) if cell_ref else "")
            if "_row" in row and isinstance(row["_row"], int):
                collected_rows.append(row["_row"])

        if len(collected_values) <= 1:
            continue

        has_conflict = False
        first_val = collected_values[0]
        for val in collected_values[1:]:
            if not _are_values_equal(first_val, val):
                has_conflict = True
                break

        if has_conflict:
            item_id = f"rev_conflict_{uuid.uuid4().hex[:12]}"
            cell_desc = " vs ".join(filter(None, collected_cells)) or f"rows {collected_rows}"
            source_ctx = f"Cells {cell_desc} for document {target_key} attribute '{attr}'"
            review_item = DocumentReviewItem(
                item_id=item_id,
                category="INTRA_WORKBOOK_CONFLICT",
                target_document_key=target_key,
                source_label=attr,
                source_context=source_ctx,
                raw_values=collected_values,
                candidate_resolutions=[],
                status="pending",
                decision=None,
                notes=f"Intra-workbook conflict detected on '{attr}' between {cell_desc}",
                run_id=run_id,
                payload={
                    "cells": collected_cells,
                    "rows": collected_rows,
                    "raw_values": collected_values,
                    "target_document_key": target_key,
                    "attribute": attr,
                },
            )
            conflicts.append(review_item)

    return conflicts


def evaluate_photo_variety_alignment(
    photo_signals: list[PhotoContextSignal | dict[str, Any]],
    plant_identities: list[PlantIdentity],
) -> list[DocumentReviewItem]:
    """Check candidate photo signals against document variety identities."""
    review_items: list[DocumentReviewItem] = []

    plant_tokens: set[str] = set()
    for plant in plant_identities:
        text_corpus = f"{plant.id} {plant.label} {plant.scientific_name or ''} {plant.source_category or ''}".lower()
        for tok in re.findall(r"[a-z0-9]+", text_corpus):
            plant_tokens.add(tok)

    stop_tokens = {"cabai", "cabe", "capsicum", "tanaman", "foto", "buah", "daun", "l", "sp"}

    candidate_resolutions = [
        {"action": "ASSIGN_EXISTING", "description": "Assign photo to an existing document variety"},
        {"action": "CREATE_NEW_DOCUMENT", "description": "Create new document for novel variety"},
        {"action": "DISCARD", "description": "Discard mismatched photo"},
    ]

    for item in photo_signals:
        signal = item if isinstance(item, PhotoContextSignal) else PhotoContextSignal(**item)

        hint = signal.suggested_variety_hint or ""
        if not hint and signal.extracted_variety_hints:
            hint = signal.extracted_variety_hints[0]

        if not hint:
            continue

        hint_tokens = set(re.findall(r"[a-z0-9]+", hint.lower())) - stop_tokens
        if not hint_tokens:
            continue

        matched = any(tok in plant_tokens for tok in hint_tokens)
        if not matched:
            rev_id = f"rev_photo_{uuid.uuid4().hex[:12]}"
            review_items.append(
                DocumentReviewItem(
                    item_id=rev_id,
                    category="PHOTO_VARIETY_MISMATCH",
                    target_document_key=None,
                    source_label=signal.filename,
                    source_context=f"Photo '{signal.filename}' suggested variety '{hint}' mismatches document varieties",
                    raw_values=[hint],
                    candidate_resolutions=candidate_resolutions,
                    status="pending",
                    decision=None,
                    notes=f"Photo variety mismatch for {signal.filename}: hint '{hint}' not found in document varieties",
                    payload={
                        "filename": signal.filename,
                        "suggested_variety_hint": hint,
                        "caption": signal.caption,
                    },
                )
            )

    return review_items


class DocumentPlanner:
    """Document staging planner and intra-workbook conflict detector."""

    def __init__(
        self,
        catalog: AttributeCatalog | None = None,
        policy: AttributePolicy | None = None,
        location_registry: AdministrativeRegistry | None = None,
        plant_catalog: PlantCatalog | None = None,
    ) -> None:
        self.catalog = catalog or AttributeCatalog()
        self.policy = policy or AttributePolicy(catalog=self.catalog)
        self.location_registry = location_registry
        self.plant_catalog = plant_catalog
        self._ensure_known_definitions()

    def _ensure_known_definitions(self) -> None:
        """Register default field measurements into catalog if missing."""
        for item in _KNOWN_FIELD_DEFINITIONS:
            if not self.catalog.lookup(item["attribute_id"]):
                defn = AttributeDefinition(**item)
                try:
                    self.catalog.register_attribute(defn)
                except ValueError:
                    pass

    def parse_cell_value(
        self,
        value: Any,
        expected_type: str | None = None,
        unit: str | None = None,
    ) -> ParsedCellValue:
        """Delegate cell parsing to module-level parse_cell_value function."""
        return parse_cell_value(value, expected_type=expected_type, unit=unit)

    def plan_records(
        self,
        records: list[dict[str, Any]],
        current_documents: dict[str, KnowledgeDocument] | None = None,
        base_versions: dict[str, int] | None = None,
        photo_signals: list[PhotoContextSignal | dict[str, Any]] | None = None,
        run_id: str = "run_default",
        source_filename: str = "data_input.xlsx",
        source_scope: str = "field_survey_2026",
        metadata: dict[str, Any] | None = None,
    ) -> DocumentDraft:
        """Group records into KnowledgeDocuments, resolve identities, and detect conflicts."""
        reviews: list[DocumentReviewItem] = []
        resolved_records_by_key: dict[
            str,
            list[tuple[PlantIdentity, HabitatIdentity, SamplingLocationDetail, dict[str, Any]]]
        ] = {}

        # 1. Resolve identity and group records by (variety.id, habitat.id)
        for row in records:
            raw_address = (
                row.get("Lokasi Sampling")
                or row.get("lokasi_sampling")
                or row.get("alamat")
                or row.get("address")
                or ""
            )
            raw_variety = (
                row.get("Varietas")
                or row.get("varietas")
                or row.get("Jenis Cabai")
                or row.get("jenis_cabai")
                or row.get("variety")
                or ""
            )

            loc_ctx: dict[str, Any] = {}
            if "S" in row and "E" in row:
                try:
                    s_val = float(row["S"])
                    e_val = float(row["E"])
                    loc_ctx["coordinates"] = {
                        "s": s_val,
                        "e": e_val,
                        "latitude": s_val,
                        "longitude": e_val,
                    }
                except (ValueError, TypeError):
                    pass
            elif "coordinates" in row and isinstance(row["coordinates"], dict):
                loc_ctx["coordinates"] = row["coordinates"]

            if "Ketinggian Lahan Sampling (mdpl = meter diatas permukaan laut)" in row:
                loc_ctx["elevation_mdpl"] = row[
                    "Ketinggian Lahan Sampling (mdpl = meter diatas permukaan laut)"
                ]
            elif "elevation" in row:
                loc_ctx["elevation_mdpl"] = row["elevation"]

            if "pH tanah" in row:
                loc_ctx["soil_ph"] = row["pH tanah"]

            habitat: HabitatIdentity | None = None
            loc_detail: SamplingLocationDetail | None = None
            try:
                habitat, loc_detail, issues = resolve_sampling_location(
                    str(raw_address),
                    context=loc_ctx,
                    registry=self.location_registry,
                )
                if habitat is None or issues:
                    reviews.append(
                        DocumentReviewItem(
                            item_id=f"rev_loc_{uuid.uuid4().hex[:12]}",
                            category="LOCATION_UNRESOLVED",
                            source_label=str(raw_address),
                            source_context=f"Row {row.get('_row', 'unknown')}: {raw_address}",
                            status="pending",
                            run_id=run_id,
                            notes=f"Unresolved administrative address: {raw_address}",
                            raw_values=[raw_address],
                        )
                    )
            except Exception as e:
                reviews.append(
                    DocumentReviewItem(
                        item_id=f"rev_loc_{uuid.uuid4().hex[:12]}",
                        category="LOCATION_UNRESOLVED",
                        source_label=str(raw_address),
                        source_context=f"Row {row.get('_row', 'unknown')}: {raw_address}",
                        status="pending",
                        run_id=run_id,
                        notes=f"Location resolution failed: {e}",
                        raw_values=[raw_address],
                    )
                )

            variety: PlantIdentity | None = None
            try:
                resolved_plant = resolve_plant_identity(
                    str(raw_variety),
                    catalog=self.plant_catalog,
                )
                if raw_variety and str(raw_variety).strip():
                    variety = resolved_plant.model_copy(
                        update={"label": str(raw_variety).strip()}
                    )
                else:
                    variety = resolved_plant
            except Exception as e:
                reviews.append(
                    DocumentReviewItem(
                        item_id=f"rev_plant_{uuid.uuid4().hex[:12]}",
                        category="PLANT_IDENTITY_UNRESOLVED",
                        source_label=str(raw_variety),
                        source_context=f"Row {row.get('_row', 'unknown')}: {raw_variety}",
                        status="pending",
                        run_id=run_id,
                        notes=f"Plant identity resolution failed: {e}",
                        raw_values=[raw_variety],
                    )
                )

            if habitat is not None and variety is not None and loc_detail is not None:
                doc_key = f"{variety.id}:{habitat.id}"
                final_variety = variety
                final_habitat = habitat

                if current_documents:
                    if doc_key in current_documents:
                        final_variety = current_documents[doc_key].variety
                        final_habitat = current_documents[doc_key].habitat
                    else:
                        for exist_key, exist_doc in current_documents.items():
                            var_match = False
                            if exist_doc.variety.id == variety.id:
                                var_match = True
                            elif (
                                exist_doc.variety.label.strip().lower() == variety.label.strip().lower()
                                or exist_doc.variety.label.strip().lower() == str(raw_variety).strip().lower()
                            ):
                                var_match = True

                            hab_match = False
                            if exist_doc.habitat.id == habitat.id:
                                hab_match = True
                            elif (
                                exist_doc.habitat.label.strip().lower() == habitat.label.strip().lower()
                                or exist_doc.habitat.label.strip().lower() == str(raw_address).strip().lower()
                            ):
                                hab_match = True
                            else:
                                ep = exist_doc.habitat.administrative_path
                                hp = habitat.administrative_path
                                if (
                                    ep.village_name.strip().lower() == hp.village_name.strip().lower()
                                    and ep.regency_name.strip().lower() == hp.regency_name.strip().lower()
                                ):
                                    hab_match = True

                            if var_match and hab_match:
                                doc_key = exist_key
                                final_variety = exist_doc.variety
                                final_habitat = exist_doc.habitat
                                break

                if doc_key not in resolved_records_by_key:
                    resolved_records_by_key[doc_key] = []
                resolved_records_by_key[doc_key].append(
                    (final_variety, final_habitat, loc_detail, row)
                )

        # 2. Check intra-workbook conflicts for each group
        conflicting_attrs_by_doc: dict[str, set[str]] = {}
        for doc_key, grouped_tuples in resolved_records_by_key.items():
            grouped_rows = [t[3] for t in grouped_tuples]
            doc_conflicts = detect_intra_workbook_conflicts(
                doc_key, grouped_rows, run_id=run_id
            )
            if doc_conflicts:
                reviews.extend(doc_conflicts)
                conflicting_attrs_by_doc[doc_key] = {
                    c.source_label for c in doc_conflicts
                }

        # 3. Construct KnowledgeDocument per grouped envelope
        documents: dict[str, KnowledgeDocument] = {}
        current_docs = current_documents or {}

        for doc_key, grouped_tuples in resolved_records_by_key.items():
            first_variety = grouped_tuples[0][0]
            first_habitat = grouped_tuples[0][1]

            if doc_key in current_docs:
                base_doc = current_docs[doc_key]
                doc_attributes = copy.deepcopy(base_doc.attributes)
                doc_sampling_locations = copy.deepcopy(base_doc.sampling_locations)
                doc_media_refs = copy.deepcopy(base_doc.media_refs)
            else:
                doc_attributes = {}
                doc_sampling_locations = []
                doc_media_refs = []

            for _, _, loc_detail, _ in grouped_tuples:
                is_duplicate = any(
                    existing.hamlet == loc_detail.hamlet
                    and existing.sampling_marker == loc_detail.sampling_marker
                    and existing.raw_address == loc_detail.raw_address
                    for existing in doc_sampling_locations
                )
                if not is_duplicate:
                    doc_sampling_locations.append(loc_detail)

            conflicts_for_doc = conflicting_attrs_by_doc.get(doc_key, set())

            ignore_keys = {
                "lokasi sampling", "lokasi_sampling", "alamat", "location", "address",
                "varietas", "jenis cabai", "jenis_cabai", "plant", "variety",
                "s", "e", "coordinates", "latitude", "longitude",
                "ketinggian lahan sampling (mdpl = meter diatas permukaan laut)",
                "no", "elevation",
            }

            for _, _, loc_detail, row in grouped_tuples:
                for k, v in row.items():
                    if k.startswith("_") or k.strip().lower() in ignore_keys:
                        continue

                    defn = self.catalog.lookup(k)
                    if defn is None:
                        eval_res = self.policy.evaluate(k, catalog=self.catalog)
                        if eval_res.status == "AUTO_ACCEPT" and eval_res.attribute_id:
                            defn = self.catalog.lookup(eval_res.attribute_id)

                    if defn is None:
                        reviews.append(
                            DocumentReviewItem(
                                item_id=f"rev_attr_{uuid.uuid4().hex[:12]}",
                                category="NEW_ATTRIBUTE",
                                source_label=k,
                                source_context=f"Row {row.get('_row', 'unknown')}: Novel attribute '{k}'",
                                raw_values=[v],
                                status="pending",
                                run_id=run_id,
                                notes=f"Novel attribute '{k}' not registered in catalog",
                            )
                        )
                        continue

                    attr_id = defn.attribute_id

                    parsed_cell = self.parse_cell_value(
                        v,
                        expected_type=defn.value_type,
                        unit=defn.unit,
                    )

                    if not parsed_cell.is_valid:
                        if (
                            "unit" in (parsed_cell.error or "").lower()
                            and attr_id in doc_attributes
                        ):
                            existing_attr = doc_attributes[attr_id]
                            reviews.append(
                                DocumentReviewItem(
                                    item_id=f"rev_ow_{uuid.uuid4().hex[:12]}",
                                    category="UNIT_MISMATCH",
                                    source_label=k,
                                    target_document_key=doc_key,
                                    source_context=f"Row {row.get('_row', 'unknown')}: Attribute '{k}' unit mismatch",
                                    raw_values=[
                                        existing_attr.value,
                                        parsed_cell.value if parsed_cell.value is not None else v,
                                    ],
                                    status="pending",
                                    run_id=run_id,
                                    notes=(
                                        f"Unit mismatch for '{k}': existing unit={existing_attr.unit} "
                                        f"vs incoming unit={parsed_cell.unit}"
                                    ),
                                    payload={
                                        "attribute_id": attr_id,
                                        "old_value": existing_attr.value,
                                        "new_value": (
                                            parsed_cell.value
                                            if parsed_cell.value is not None
                                            else v
                                        ),
                                        "old_unit": existing_attr.unit,
                                        "new_unit": parsed_cell.unit,
                                        "old_attribute": existing_attr.model_dump(mode="json"),
                                        "raw_value": v,
                                    },
                                )
                            )
                        else:
                            reviews.append(
                                DocumentReviewItem(
                                    item_id=f"rev_val_{uuid.uuid4().hex[:12]}",
                                    category="VALUE_INVALID",
                                    source_label=k,
                                    source_context=f"Row {row.get('_row', 'unknown')}: Invalid value '{v}' for attribute '{k}'",
                                    raw_values=[v],
                                    status="pending",
                                    run_id=run_id,
                                    notes=parsed_cell.error or "Invalid attribute value or unit mismatch",
                                )
                            )
                        continue

                    if parsed_cell.is_blank:
                        continue

                    if k in conflicts_for_doc or attr_id in conflicts_for_doc:
                        if attr_id in doc_attributes:
                            del doc_attributes[attr_id]
                        continue

                    cell_ref = (
                        row.get(f"_cell_{k}")
                        or row.get("_cell")
                        or (f"row_{row.get('_row')}" if row.get("_row") else "A1")
                    )
                    row_idx = row.get("_row", 1) if isinstance(row.get("_row"), int) else 1
                    sheet_name = str(row.get("_sheet", "Resume Data"))

                    incoming_attr = AttributeValue(
                        label=defn.label,
                        value=parsed_cell.value,
                        value_type=parsed_cell.value_type or defn.value_type,
                        unit=parsed_cell.unit or defn.unit,
                        resolution="AUTO_ACCEPT",
                        mapping_policy_version="exact-alias-policy-v1",
                        current_source=CurrentSourceReference(
                            source_ref=source_filename,
                            sheet=sheet_name,
                            row=row_idx,
                            cell=str(cell_ref),
                            location_ref=loc_detail.location_ref,
                        ),
                    )

                    if attr_id in doc_attributes:
                        existing_attr = doc_attributes[attr_id]
                        if not self.policy.can_overwrite(existing_attr, incoming_attr):
                            category = (
                                "UNIT_MISMATCH"
                                if (
                                    existing_attr.unit is not None
                                    and incoming_attr.unit is not None
                                    and not self.policy._units_compatible(
                                        existing_attr.unit, incoming_attr.unit
                                    )
                                )
                                else "OVERWRITE_BLOCKED"
                            )
                            reviews.append(
                                DocumentReviewItem(
                                    item_id=f"rev_ow_{uuid.uuid4().hex[:12]}",
                                    category=category,
                                    source_label=k,
                                    target_document_key=doc_key,
                                    source_context=f"Row {row.get('_row', 'unknown')}: Attribute '{k}' overwrite blocked",
                                    raw_values=[existing_attr.value, incoming_attr.value],
                                    status="pending",
                                    run_id=run_id,
                                    notes=(
                                        f"Overwrite blocked for '{k}': existing value={existing_attr.value} "
                                        f"(type={existing_attr.value_type}, unit={existing_attr.unit}) vs "
                                        f"incoming value={incoming_attr.value} "
                                        f"(type={incoming_attr.value_type}, unit={incoming_attr.unit})"
                                    ),
                                    payload={
                                        "attribute_id": attr_id,
                                        "old_value": existing_attr.value,
                                        "new_value": incoming_attr.value,
                                        "old_attribute": existing_attr.model_dump(mode="json"),
                                        "new_attribute": incoming_attr.model_dump(mode="json"),
                                    },
                                )
                            )
                            continue

                    doc_attributes[attr_id] = incoming_attr

            knowledge_doc = KnowledgeDocument(
                schema_version="kms-document-v1",
                variety=first_variety,
                habitat=first_habitat,
                sampling_locations=doc_sampling_locations,
                attributes=doc_attributes,
                media_refs=doc_media_refs,
                updated_by_run=run_id,
            )
            documents[doc_key] = knowledge_doc

        # 4. Check photo signals variety alignment (A24) and link matching media
        if photo_signals:
            all_plants = [doc.variety for doc in documents.values()]
            photo_reviews = evaluate_photo_variety_alignment(photo_signals, all_plants)
            reviews.extend(photo_reviews)

            stop_tokens = {"cabai", "cabe", "capsicum", "tanaman", "foto", "buah", "daun", "l", "sp"}
            for item in photo_signals:
                signal = item if isinstance(item, PhotoContextSignal) else PhotoContextSignal(**item)
                hint = signal.suggested_variety_hint or ""
                if not hint and signal.extracted_variety_hints:
                    hint = signal.extracted_variety_hints[0]
                if not hint:
                    continue

                hint_tokens = set(re.findall(r"[a-z0-9]+", hint.lower())) - stop_tokens
                if not hint_tokens:
                    continue

                for doc in documents.values():
                    plant = doc.variety
                    text_corpus = f"{plant.id} {plant.label} {plant.scientific_name or ''} {plant.source_category or ''}".lower()
                    p_tokens = set(re.findall(r"[a-z0-9]+", text_corpus))
                    if any(tok in p_tokens for tok in hint_tokens):
                        organ_val = (signal.organ or "daun").lower()
                        if organ_val not in ("habitus", "daun", "batang", "bunga", "buah", "biji"):
                            organ_val = "daun"
                        asset_id = hashlib.sha256(signal.filename.encode("utf-8")).hexdigest()
                        if not any(m.asset_id == asset_id for m in doc.media_refs):
                            doc.media_refs.append(
                                MediaReference(
                                    asset_id=asset_id,
                                    organ_part=organ_val,
                                    caption=signal.caption,
                                    source_filename=signal.filename,
                                    provenance_run_id=run_id,
                                    assigned_via="AUTO_ACCEPT",
                                    resolution_status="linked",
                                )
                            )

        # 5. Populate base versions
        calculated_base_versions: dict[str, int] = {}
        if base_versions:
            calculated_base_versions.update(base_versions)
        elif current_documents:
            for k, doc in current_documents.items():
                calculated_base_versions[k] = doc.version

        # 6. Build file and template hashes
        template_file = Path("data/canonical/template_kanonik.xlsx")
        if template_file.exists():
            tmpl_hash = hashlib.sha256(template_file.read_bytes()).hexdigest()
        else:
            tmpl_hash = hashlib.sha256(b"canonical_template_v1").hexdigest()

        file_hash = hashlib.sha256(source_filename.encode("utf-8")).hexdigest()

        return DocumentDraft(
            run_id=run_id,
            source_filename=source_filename,
            file_sha256=file_hash,
            source_uri=None,
            source_scope=source_scope,
            schema_version="kms-document-v1",
            template_hash=tmpl_hash,
            config_fingerprint="kms-planner-v1",
            revision=1,
            documents=documents,
            reviews=reviews,
            base_versions=calculated_base_versions,
            metadata=metadata or {},
        )

    plan_document_draft = plan_records

    def plan_workbook(
        self,
        workbook_path: Path | str,
        run_id: str = "run_default",
        current_documents: dict[str, KnowledgeDocument] | None = None,
        base_versions: dict[str, int] | None = None,
        photo_signals: list[PhotoContextSignal | dict[str, Any]] | None = None,
        source_scope: str = "field_survey_2026",
        metadata: dict[str, Any] | None = None,
    ) -> DocumentDraft:
        """Parse real field workbook (data_input.xlsx) and plan document envelopes."""
        path = Path(workbook_path)
        if not path.exists():
            raise FileNotFoundError(f"Workbook not found at: {path}")

        file_bytes = path.read_bytes()
        calculated_file_sha = hashlib.sha256(file_bytes).hexdigest()

        wb = openpyxl.load_workbook(path, data_only=True)
        try:
            if "Resume Data" in wb.sheetnames:
                ws = wb["Resume Data"]
            else:
                ws = wb.active

            col_headers: dict[int, str] = {}
            for col_idx in range(1, ws.max_column + 1):
                val1 = ws.cell(row=1, column=col_idx).value
                val2 = ws.cell(row=2, column=col_idx).value if ws.max_row >= 2 else None

                if val2 is not None and str(val2).strip():
                    col_headers[col_idx] = str(val2).strip()
                elif val1 is not None and str(val1).strip():
                    col_headers[col_idx] = str(val1).strip()

            records: list[dict[str, Any]] = []
            for row_idx in range(3, ws.max_row + 1):
                has_content = any(
                    ws.cell(row=row_idx, column=c).value is not None
                    for c in col_headers.keys()
                )
                if not has_content:
                    continue

                rec: dict[str, Any] = {"_row": row_idx, "_sheet": ws.title}
                for col_idx, header_name in col_headers.items():
                    cell = ws.cell(row=row_idx, column=col_idx)
                    rec[header_name] = cell.value
                    col_letter = openpyxl.utils.get_column_letter(col_idx)
                    rec[f"_cell_{header_name}"] = f"{col_letter}{row_idx}"

                if "Jenis Cabai" in rec and "Varietas" not in rec:
                    rec["Varietas"] = rec["Jenis Cabai"]

                records.append(rec)
        finally:
            wb.close()

        draft = self.plan_records(
            records=records,
            current_documents=current_documents,
            base_versions=base_versions,
            photo_signals=photo_signals,
            run_id=run_id,
            source_filename=path.name,
            source_scope=source_scope,
            metadata=metadata,
        )

        draft = draft.model_copy(update={"file_sha256": calculated_file_sha})
        return draft


def plan_document_changes(
    source_records: list[dict[str, Any]],
    current_documents: dict[str, KnowledgeDocument] | None = None,
    run_id: str = "run_default",
    **kwargs: Any,
) -> DocumentDraft:
    """Top-level functional entrypoint matching DocumentPlanner."""
    planner = DocumentPlanner()
    return planner.plan_records(
        records=source_records,
        current_documents=current_documents,
        run_id=run_id,
        **kwargs,
    )


plan_document_draft = plan_document_changes
