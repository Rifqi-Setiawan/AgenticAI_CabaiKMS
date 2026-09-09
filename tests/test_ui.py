from __future__ import annotations

import io
from unittest.mock import patch

import openpyxl
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from src.ui.pipeline_runner import PipelineRunResult
from src.agents.schema_matching import review_queue
from src.schema.canonical import CanonicalSchema
from src.schema.contracts import NULL_ROW, SchemaMapping
from src.ui.output_builder import CanonicalOutputBuilder, worksheet_to_dataframe
from src.ui.pipeline_runner import _deterministic_workbook_bytes

APP_PATH = "src/ui/app.py"
PROGRESS_PATH = "src/ui/pages/2_Progress.py"
HASIL_PATH = "src/ui/pages/3_Hasil.py"

# app.py's import chain now pulls in the full pipeline stack (torch,
# transformers, langgraph, chromadb, google api client, ...) — the
# default 3s AppTest timeout is comfortably enough once warm, but a
# cold run reliably exceeds it.
APP_TEST_TIMEOUT = 30


def _fixture_result(with_issues: bool = False) -> PipelineRunResult:
    mapping_df = pd.DataFrame(
        [
            {
                "source_attribute_display": "Warna Daun",
                "source_attribute": "Warna Daun",
                "source_context": None,
                "predicted_row": "r_7",
                "predicted_label": "warna daun",
                "proposed_target_canonical_key": "daun.warna",
                "target_domain": "daun",
                "mapping_method": "exact_name",
                "confidence": 0.9,
                "verifier_status": "PASS",
                "acceptance_status": "AUTO_ACCEPT",
                "canonical_write": True,
                "normalization_required": False,
                "reasoning": "cocok jelas dengan label 'warna daun'",
            },
            {
                "source_attribute_display": "Panjang Daun",
                "source_attribute": "Panjang Daun",
                "source_context": None,
                "predicted_row": "r_8",
                "predicted_label": "panjang daun",
                "proposed_target_canonical_key": "daun.panjang",
                "target_domain": "daun",
                "mapping_method": "retrieve_rerank",
                "confidence": 0.85,
                "verifier_status": "REVIEW",
                "acceptance_status": "REVIEW",
                "canonical_write": False,
                "normalization_required": True,
                "reasoning": "cocok dengan tipe data numerik",
            },
        ]
    )
    canonical_df = pd.DataFrame(
        [
            {"Nomor": 1, "Karakter": "habitus", "Gendot": "perdu", "Kopay": "terna"},
            {"Nomor": 7, "Karakter": "warna daun", "Gendot": "green group 137 A", "Kopay": ""},
            {"Nomor": 8, "Karakter": "panjang daun", "Gendot": "10--12 cm", "Kopay": "8--10 cm"},
        ]
    )
    vision_rows = [
        {
            "filename": "daun1.jpg", "status": "KNOWN", "matched_variety": "Gendot",
            "identified_part": "DAUN", "confidence": 0.9, "visual_evidence": "ok",
            "write_applied": False, "write_reason": "target row not found",
        }
    ]
    error_trace = ["contoh entri error_trace"] if with_issues else []
    return PipelineRunResult(
        mapping_df=mapping_df,
        canonical_df=canonical_df,
        workbook_bytes=b"fake-xlsx-bytes-for-ui-tests",
        vision_rows=vision_rows,
        agent_status={
            "schema_matching": "selesai — 2 dipetakan, 0 review",
            "vision_classification": "selesai — 1 citra diklasifikasi, 0 UNCERTAIN",
            "orchestrator": "checkpoint tersimpan (thread_id=test-thread-fixture)",
        },
        checkpoint_thread_id="test-thread-fixture-nonexistent",
        error_trace=error_trace,
    )


def _xlsx_bytes(*sheet_names: str) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_names[0]
    for name in sheet_names[1:]:
        workbook.create_sheet(name)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _real_null_review_result(tmp_path) -> tuple[PipelineRunResult, str]:
    schema = CanonicalSchema.from_template()
    queue_path = tmp_path / "review.jsonl"
    run_id = "ui-replay-run"
    source_hash = "b" * 64
    mapping = SchemaMapping(
        source_attribute="Growth habit", source_context="Vegetative",
        source_format="row-oriented", target_canonical_row=NULL_ROW,
        confidence=0.2, reasoning="no confident model proposal",
        normalization_required=True,
    )
    item = review_queue.enqueue(
        mapping, reason="needs review", queue_path=queue_path,
        run_id=run_id, mapping_item_id="mapping-ui-replay",
        source_file_name="source.xlsx", source_file_sha256=source_hash,
        source_sheet="Observations", source_format="row-oriented",
        source_attribute_display="Vegetative / Growth habit",
        schema_version=schema.schema_version, template_hash=schema.template_hash,
        proposed_canonical_key=None, mapping_method="retrieve_rerank",
        verifier_status="REVIEW", sample_values=["terna"],
        raw_values_by_variety={"Domba": ["terna"]},
    )
    builder = CanonicalOutputBuilder(schema=schema, variety_names=["Domba"])
    workbook = builder.build_workbook()
    try:
        workbook_bytes = _deterministic_workbook_bytes(workbook)
        canonical_df = worksheet_to_dataframe(workbook["Sheet1"], schema, ["Domba"])
    finally:
        workbook.close()
    mapping_df = pd.DataFrame([{
        "review_item_id": item.item_id,
        "mapping_item_id": "mapping-ui-replay",
        "source_file_sha256": source_hash,
        "source_sheet": "Observations",
        "source_format": "row-oriented",
        "source_attribute_display": "Vegetative / Growth habit",
        "source_attribute": "Growth habit",
        "source_context": "Vegetative",
        "predicted_row": NULL_ROW,
        "predicted_label": "NULL",
        "target_domain": None,
        "confidence": 0.2,
        "normalization_required": True,
        "reasoning": "no confident model proposal",
        "mapping_method": "retrieve_rerank",
        "verifier_status": "REVIEW",
        "acceptance_status": "REVIEW",
        "canonical_write": False,
        "proposed_target_canonical_key": None,
    }])
    return PipelineRunResult(
        mapping_df=mapping_df, canonical_df=canonical_df,
        workbook_bytes=workbook_bytes, vision_rows=[], run_id=run_id,
        checkpoint_thread_id=run_id, review_queue_path=str(queue_path),
        schema_version=schema.schema_version, template_hash=schema.template_hash,
    ), schema.row_by_label("habitus").canonical_key


class TestPage1Input:
    def test_uploader_accepts_xlsx_only(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT).run()
        uploader = at.get("file_uploader")[0]
        assert uploader.allowed_type == [".xlsx"]

    def test_header_selection_only_for_row_oriented(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT).run()
        assert at.selectbox[0].options == ["Otomatis", "1 baris", "2 baris"]
        at.selectbox[0].select("2 baris").run()
        assert not at.exception
        at.radio[0].set_value("transposed").run()
        assert not at.selectbox

    def test_renders_without_exception(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.run()
        assert not at.exception

    def test_has_upload_and_drive_inputs_and_run_button(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.run()
        assert len(at.text_input) >= 1
        assert len(at.radio) >= 1
        assert any("Jalankan Pipeline" in b.label for b in at.button)

    def test_run_without_file_is_disabled(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.run()
        run_button = next(b for b in at.button if "Jalankan Pipeline" in b.label)
        assert run_button.disabled

    def test_invalid_workbook_preflight_blocks_run(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT).run()
        at.get("file_uploader")[0].upload("broken.xlsx", b"not an xlsx").run()
        assert not at.exception
        assert any("tidak dapat dibuka" in error.value for error in at.error)
        assert next(b for b in at.button if "Jalankan Pipeline" in b.label).disabled

    def test_multisheet_selection_and_max_images_reach_runtime(self):
        calls = []

        def fake_run(file_path, **kwargs):
            calls.append(kwargs)
            return _fixture_result()

        with patch("src.ui.pipeline_runner.run_pipeline_ui", fake_run):
            at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT).run()
            at.get("file_uploader")[0].upload(
                "demo.xlsx", _xlsx_bytes("First", "Chosen")
            ).run()
            sheet = next(box for box in at.selectbox if box.label == "Worksheet yang diproses")
            sheet.select("Chosen").run()
            max_images = next(item for item in at.number_input if item.label == "Maximum images to process")
            max_images.set_value(7).run()
            next(b for b in at.button if "Jalankan Pipeline" in b.label).click().run()

        assert not at.exception
        assert calls[-1]["sheet_name"] == "Chosen"
        assert calls[-1]["max_images"] == 7
        assert calls[-1]["drive_folder_id"] == ""
        assert len(calls) == 1

    def test_shows_last_result_hint_when_result_exists(self):
        at = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result()
        at.session_state["cabai_kms_inputs"] = {"filename": "contoh.xlsx"}
        at.run()
        assert not at.exception
        assert any("contoh.xlsx" in i.value for i in at.info)


class TestPage2Progress:
    def test_empty_state_shows_info(self):
        at = AppTest.from_file(PROGRESS_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.run()
        assert not at.exception
        assert any("Belum ada pipeline" in i.value for i in at.info)

    def test_with_result_shows_agent_metrics_and_log(self):
        at = AppTest.from_file(PROGRESS_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result()
        at.session_state["cabai_kms_log"] = ["mulai...", "schema_matching: ...", "selesai."]
        at.run()
        assert not at.exception
        metric_labels = [m.label for m in at.metric]
        assert "schema_matching/tabular" in metric_labels
        assert "vision_classification" in metric_labels
        assert "drive_crawler" in metric_labels
        assert any("selesai." in c.value for c in at.code)

    def test_no_drive_is_skipped_not_failed(self):
        result = _fixture_result()
        result.agent_status["drive_crawler"] = "dilewati (tidak ada folder Drive)"
        result.agent_status["vision_classification"] = "dilewati (tidak ada folder Drive)"
        at = AppTest.from_file(PROGRESS_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = result
        at.run()
        values = {metric.label: metric.value for metric in at.metric}
        assert values["drive_crawler"] == "SKIPPED"
        assert values["vision_classification"] == "SKIPPED"

    def test_zero_review_and_zero_uncertain_are_success(self):
        result = _fixture_result()
        result.agent_status["schema_matching"] = (
            "selesai — 13 atribut: 13 AUTO_ACCEPT, 0 REVIEW, 0 NO_WRITE"
        )
        result.agent_status["vision_classification"] = (
            "selesai — 5 citra diklasifikasi, 5 ditulis ke sel, 0 UNCERTAIN"
        )
        at = AppTest.from_file(PROGRESS_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = result
        at.run()
        values = {metric.label: metric.value for metric in at.metric}
        assert values["schema_matching/tabular"] == "SUCCESS"
        assert values["vision_classification"] == "SUCCESS"

    def test_no_issues_shows_success(self):
        at = AppTest.from_file(PROGRESS_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result(with_issues=False)
        at.run()
        assert not at.exception
        assert any("Tidak ada" in s.value for s in at.success)

    def test_with_issues_shows_warning(self):
        at = AppTest.from_file(PROGRESS_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result(with_issues=True)
        at.run()
        assert not at.exception
        assert any("catatan runtime/peringatan" in w.value for w in at.warning)
        assert all("manual_review" not in w.value for w in at.warning)


class TestPage3Hasil:
    def test_empty_state_shows_info_and_stops(self):
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.run()
        assert not at.exception
        assert any("Belum ada pipeline" in i.value for i in at.info)
        assert len(at.dataframe) == 0  # st.stop() reached, nothing below rendered

    def test_with_result_shows_tables_and_download_button(self):
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result()
        at.run()
        assert not at.exception
        assert len(at.dataframe) >= 2  # canonical + mapping (+ vision)
        assert any("Unduh hasil" in b.label for b in list(at.button) + list(at.download_button))

    def test_summary_and_mapping_inspector_use_actual_decision_fields(self):
        result = _fixture_result()
        third = result.mapping_df.iloc[[1]].copy()
        third["source_attribute_display"] = "Unknown"
        third["acceptance_status"] = "NO_WRITE"
        third["canonical_write"] = False
        result.mapping_df = pd.concat([result.mapping_df, third], ignore_index=True)
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = result
        at.session_state["cabai_kms_inputs"] = {"filename": "demo.xlsx", "sheet_name": "Data", "drive_url": ""}
        at.run()
        assert not at.exception
        metrics = {metric.label: metric.value for metric in at.metric}
        assert metrics["AUTO_ACCEPT"] == "1"
        assert metrics["REVIEW"] == "1"
        assert metrics["NO_WRITE"] == "1"
        mapping_columns = set(at.dataframe[1].value.columns)
        assert {"Proposed canonical label", "Mapping method", "Verifier status", "Acceptance status", "Canonical write"} <= mapping_columns
        assert any("Prediksi model tidak sama" in caption.value for caption in at.caption)

    def test_no_drive_reports_vision_skipped_without_error(self):
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result()
        at.session_state["cabai_kms_inputs"] = {"filename": "demo.xlsx", "sheet_name": "Data", "drive_url": ""}
        at.run()
        assert any("Vision skipped" in info.value for info in at.info)
        assert not at.error

    def test_vision_view_exposes_write_reason(self):
        result = _fixture_result()
        result.images_discovered = 1
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = result
        at.session_state["cabai_kms_inputs"] = {"filename": "demo.xlsx", "sheet_name": "Data", "drive_url": "folder"}
        at.run()
        assert not at.exception
        assert "write_reason" in at.dataframe[2].value.columns
        assert at.dataframe[2].value.iloc[0]["write_reason"] == "target row not found"
        assert any(expander.label == "Advanced / Debug" for expander in at.expander)

    def test_reasoning_selectbox_shows_full_reasoning_on_selection(self):
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result()
        at.run()
        select = at.selectbox[0]
        select.select("Panjang Daun").run()
        assert not at.exception
        assert any("cocok dengan tipe data numerik" in md.value for md in at.markdown)

    def test_checkpoint_debugger_button_does_not_crash_on_unknown_thread(self):
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = _fixture_result()
        at.run()
        debugger_button = next(b for b in at.button if "Debugger Checkpoint" in b.label)
        debugger_button.click().run()
        assert not at.exception
        assert any("checkpoint tersimpan" in w.value for w in list(at.markdown) + list(at.text))

    def test_current_run_review_item_shows_human_actions_and_canonical_selector(self, tmp_path):
        schema = CanonicalSchema.from_template()
        result = _fixture_result()
        result.run_id = "ui-current-run"
        result.schema_version = schema.schema_version
        result.template_hash = schema.template_hash
        result.review_queue_path = str(tmp_path / "review.jsonl")
        mapping = SchemaMapping(
            source_attribute="Warna Daun", source_context="Daun",
            source_format="row-oriented", target_canonical_row=schema.row_by_label("warna daun").id,
            confidence=0.2, reasoning="needs review", normalization_required=False,
        )
        item = review_queue.enqueue(
            mapping, reason="needs review", queue_path=result.review_queue_path,
            run_id=result.run_id, mapping_item_id="mapping-ui",
            source_file_name="source.xlsx", source_file_sha256="a" * 64,
            source_sheet="Sheet1", source_format="row-oriented",
            source_attribute_display="Daun / Warna Daun",
            schema_version=schema.schema_version, template_hash=schema.template_hash,
            proposed_canonical_key=schema.row_by_label("warna daun").canonical_key,
            mapping_method="retrieve_rerank", verifier_status="REVIEW",
            sample_values=["hijau"], raw_values_by_variety={"Gendot": ["hijau"]},
        )
        result.mapping_df["review_item_id"] = [item.item_id, None]
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = result
        at.run()
        assert not at.exception
        labels = [button.label for button in at.button]
        assert "Setujui usulan" in labels
        assert "Ubah target" in labels
        assert "Tandai NO_MATCH" in labels
        assert any("warna daun" in str(option) for option in at.selectbox[0].options)

    def test_revise_and_apply_correction_end_to_end(self, tmp_path):
        result, replacement_key = _real_null_review_result(tmp_path)
        at = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        at.session_state["cabai_kms_pipeline_result"] = result
        at.run()

        assert not at.exception
        assert any(
            "Vegetative / Growth habit" in expander.label
            for expander in at.expander
        )
        assert len(review_queue.list_pending(
            queue_path=result.review_queue_path, run_id=result.run_id,
        )) == 1
        at.selectbox[0].select(replacement_key).run()
        next(button for button in at.button if button.label == "Ubah target").click().run()

        assert not at.exception
        assert review_queue.list_pending(
            queue_path=result.review_queue_path, run_id=result.run_id,
        ) == []
        next(button for button in at.button if button.label == "Terapkan Koreksi").click().run()

        assert not at.exception
        corrected = at.session_state["cabai_kms_corrected_pipeline_result"]
        workbook = openpyxl.load_workbook(io.BytesIO(corrected.workbook_bytes), data_only=True)
        try:
            assert "Sheet1" in workbook.sheetnames
            assert workbook["Sheet1"]["C1"].value == "Domba"
        finally:
            workbook.close()
        habitus = corrected.canonical_df.loc[
            corrected.canonical_df["Karakter"] == "habitus", "Domba"
        ].item()
        assert habitus == "terna"
        assert any(
            button.label == "⬇️ Unduh output terkoreksi (direkomendasikan)"
            for button in at.download_button
        )
        assert at.download_button[0].label == "⬇️ Unduh output terkoreksi (direkomendasikan)"


class TestDemoScenarios:
    def test_upload_select_run_review_apply_and_download(self, tmp_path):
        result, replacement_key = _real_null_review_result(tmp_path)
        calls = []

        def fake_run(file_path, **kwargs):
            calls.append(kwargs)
            return result

        with patch("src.ui.pipeline_runner.run_pipeline_ui", fake_run):
            input_page = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT).run()
            input_page.get("file_uploader")[0].upload(
                "field demo.xlsx", _xlsx_bytes("Raw", "Demo")
            ).run()
            next(box for box in input_page.selectbox if box.label == "Worksheet yang diproses").select("Demo").run()
            next(button for button in input_page.button if "Jalankan Pipeline" in button.label).click().run()

        assert len(calls) == 1
        assert calls[0]["sheet_name"] == "Demo"
        hasil = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        hasil.session_state["cabai_kms_pipeline_result"] = input_page.session_state["cabai_kms_pipeline_result"]
        hasil.session_state["cabai_kms_inputs"] = input_page.session_state["cabai_kms_inputs"]
        hasil.run()
        assert any("field demo.xlsx" in caption.value for caption in hasil.caption)
        assert len(hasil.dataframe) >= 2
        hasil.selectbox[0].select(replacement_key).run()
        next(button for button in hasil.button if button.label == "Ubah target").click().run()
        next(button for button in hasil.button if button.label == "Terapkan Koreksi").click().run()
        assert hasil.download_button[0].label == "⬇️ Unduh output terkoreksi (direkomendasikan)"

    def test_mocked_multimodal_demo_flow_shows_write_status_and_reason(self):
        result = _fixture_result()
        result.images_discovered = 1

        with patch("src.ui.pipeline_runner.run_pipeline_ui", return_value=result):
            input_page = AppTest.from_file(APP_PATH, default_timeout=APP_TEST_TIMEOUT).run()
            input_page.get("file_uploader")[0].upload("vision.xlsx", _xlsx_bytes("Data")).run()
            input_page.text_input[0].set_value("folder-demo").run()
            next(button for button in input_page.button if "Jalankan Pipeline" in button.label).click().run()

        hasil = AppTest.from_file(HASIL_PATH, default_timeout=APP_TEST_TIMEOUT)
        hasil.session_state["cabai_kms_pipeline_result"] = input_page.session_state["cabai_kms_pipeline_result"]
        hasil.session_state["cabai_kms_inputs"] = input_page.session_state["cabai_kms_inputs"]
        hasil.run()
        vision = hasil.dataframe[2].value
        assert vision.iloc[0]["classification_status"] == "KNOWN"
        assert vision.iloc[0]["write_applied"] == False
        assert vision.iloc[0]["write_reason"] == "target row not found"
