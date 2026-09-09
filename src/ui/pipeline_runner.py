"""Fase 8 — pipeline runner behind the Streamlit UI.

Wires together what's already built (Fase 3 schema-matching, Fase 4 Drive
crawler, Fase 5 vision classification, Fase 6 tabular update, Fase 7
reliability wrappers, and the checkpointed LangGraph runtime) into one
function the UI calls. No new agent logic lives here — this module is glue
+ presentation-shaping, same spirit as eval/review_schema_matching.py.

The "hasil akhir" this produces is a real instance of the canonical
template's shape (data/canonical/template_kanonik.xlsx) — same row
labels, read dynamically — with varietas columns taken from whatever the
uploaded file's anchor column (row-oriented) or column headers
(transposed) actually contain. A canonical row nothing mapped to is left
blank; if Drive has no images (or none is given), the Gambar rows are
simply left blank too — see src/ui/output_builder.py.

Kept deliberately synchronous and capped (see max_images) — this is the
prototype phase named in the brief ("Overhead minimal; semua render dari
Python"), not a production job queue.
"""

from __future__ import annotations

import io
import re
import uuid
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Literal

import openpyxl
import pandas as pd

from src.agents.drive_crawler import DriveCrawlerError, list_images, normalize_folder_id
from src.agents.schema_matching.anchor import detect_anchor
from src.agents.schema_matching.exact_match import (
    ExactNameStatus,
    mapping_from_exact_resolution,
    resolve_exact_name,
)
from src.agents.schema_matching.exact_retrieval import build_exact_index
from src.agents.schema_matching.indexing import EMBEDDING_MODEL_NAME, ensure_indexed
from src.agents.schema_matching.mapping_verifier import (
    combine_mapping_acceptance,
    verify_mapping,
)
from src.agents.schema_matching.normalize import normalize
from src.agents.schema_matching import review_queue
from src.agents.schema_matching.review_queue import AcceptanceStatus, decide_mapping_acceptance
from src.agents.schema_matching.retrieval import (
    DEFAULT_K,
    RetrievalBackend,
    SourceAttributeProfile,
    retrieve,
    validate_retrieval_backend,
)
from src.agents.tabular_update import apply_vision_result_to_worksheet
from src.agents.vision_classification import VisionSession
from src.ingestion.runtime_source import (
    group_attribute_contributions_by_variety,
    physical_source_cells,
    prepare_legacy_runtime_source,
)
from src.ingestion.shadow_pipeline import run_structure_shadow, sanitize_shadow_error_message
from src.ingestion.source_migration import prepare_gated_runtime_source
from src.orchestrator.graph import run_pipeline  # compatibility seam for older injected tests
from src.reliability.rate_limit import (
    RateLimiter,
    RuntimeRateLimitConfig,
    describe_rate_limiter,
)
from src.reliability.wrappers import safe_classify_image, safe_rerank
from src.schema.canonical import CanonicalSchema
from src.schema.contracts import NULL_ROW
from src.schema.evaluation_config import EvaluationRunConfig
from src.schema.gold_mapping import build_mapping_item_identities
from src.schema.mapping_verification import (
    MAPPING_VERIFICATION_VERSION,
    MappingVerificationResult,
    MappingVerificationStatus,
)
from src.schema.provenance import CellProvenanceRecord, MappingMethod, source_file_sha256
from src.schema.shadow_parity import ShadowParityReport, ShadowStatus
from src.ui.output_builder import SHEET_NAME, CanonicalOutputBuilder, combine_multi_value, worksheet_to_dataframe

MAPPING_COLUMNS = [
    "mapping_item_id",
    "review_item_id",
    "mapping_identity_kind",
    "mapping_identity_value",
    "mapping_identity_issue",
    "source_file_name",
    "source_file_sha256",
    "source_sheet",
    "source_format",
    "source_attribute_display",
    "source_attribute",
    "source_context",
    "predicted_row",
    "proposed_target_canonical_key",
    "predicted_label",
    "target_domain",
    "confidence",
    "normalization_required",
    "reasoning",
    "mapping_method",
    "exact_name_status",
    "exact_name_candidates",
    "verifier_status",
    "verifier_hard_issues",
    "verifier_warnings",
    "retrieval_target_rank",
    "retrieval_target_distance",
    "retrieval_top1_row",
    "retrieval_top1_distance",
    "retrieval_top2_row",
    "retrieval_top2_distance",
    "retrieval_top1_top2_margin",
    "retrieval_target_vs_top1_gap",
    "acceptance_status",
    "acceptance_reason",
    "canonical_write",
]

ProgressCallback = Callable[[str], None]


@dataclass
class PipelineRunResult:
    mapping_df: pd.DataFrame
    canonical_df: pd.DataFrame
    workbook_bytes: bytes
    vision_rows: list[dict]
    images_discovered: int = 0
    provenance_records: list[CellProvenanceRecord] = field(default_factory=list)
    agent_status: dict[str, str] = field(default_factory=dict)
    checkpoint_thread_id: str | None = None
    error_trace: list[str] = field(default_factory=list)
    structure_shadow: ShadowParityReport | None = None
    source_backend: str = "legacy"
    source_ir_version: str | None = None
    retrieval_backend: str = "chroma"
    mapping_verifications: list[MappingVerificationResult] = field(default_factory=list)
    schema_version: str = ""
    template_hash: str = ""
    retrieval_k: int = DEFAULT_K
    mapping_verification_version: str = MAPPING_VERIFICATION_VERSION
    embedding_model_name: str | None = EMBEDDING_MODEL_NAME
    evaluation_config_fingerprint: str = ""
    run_id: str = ""
    review_queue_path: str = ""


def _noop(_: str) -> None:
    return None


def _deterministic_workbook_bytes(workbook) -> bytes:
    """Serialize equivalent workbooks identically despite OpenPyXL save timestamps."""
    raw = io.BytesIO()
    workbook.save(raw)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue()), "r") as source:
        with zipfile.ZipFile(output, "w") as target:
            target.comment = source.comment
            for original in source.infolist():
                payload = source.read(original.filename)
                if original.filename == "docProps/core.xml":
                    payload = re.sub(
                        rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",
                        rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                        payload,
                    )
                stable = zipfile.ZipInfo(original.filename, (2000, 1, 1, 0, 0, 0))
                stable.compress_type = original.compress_type
                stable.comment = original.comment
                stable.extra = original.extra
                stable.internal_attr = original.internal_attr
                stable.external_attr = original.external_attr
                stable.create_system = original.create_system
                target.writestr(stable, payload)
    return output.getvalue()


def _run_tabular_stage_impl(
    file_path: Path,
    *,
    source_format: str = "row-oriented",
    sheet_name: str | None = None,
    header_rows: int | None = None,
    k: int = DEFAULT_K,
    on_progress: ProgressCallback = _noop,
    enable_structure_shadow: bool = False,
    structure_llm_call: Callable | None = None,
    source_backend: Literal["legacy", "source-ir-gated"] = "legacy",
    retrieval_backend: RetrievalBackend = "chroma",
    embedding_encode_call: Callable[..., object] | None = None,
    review_queue_path: Path | str = review_queue.DEFAULT_QUEUE_PATH,
    text_rate_limiter: RateLimiter | None = None,
    _prepared_source_bundle=None,
    _run_id: str | None = None,
) -> PipelineRunResult:
    if source_backend not in {"legacy", "source-ir-gated"}:
        raise ValueError(
            f"unknown source backend {source_backend!r}; expected one of: "
            "legacy, source-ir-gated"
        )
    validate_retrieval_backend(retrieval_backend)

    run_id = _run_id or uuid.uuid4().hex
    resolved_sheet_name = sheet_name or _first_sheet(file_path)
    source_hash = source_file_sha256(file_path)
    state: dict = {"error_trace": []}
    agent_status: dict[str, str] = {}
    agent_status["text_rate_limit"] = describe_rate_limiter(text_rate_limiter)
    on_progress(f"text rate limit: {agent_status['text_rate_limit']}")
    structure_shadow: ShadowParityReport | None = None

    # --- backend-neutral parsing + variety-position preparation ---
    on_progress(f"Memuat berkas: {file_path.name} (format={source_format!r})")
    if _prepared_source_bundle is not None:
        source_bundle = _prepared_source_bundle
        structure_shadow = source_bundle.migration_report
        agent_status["source_ingestion"] = (
            "legacy — authoritative parser"
            if source_backend == "legacy"
            else "source-ir-gated — promoted after MATCH parity"
        )
        if structure_shadow is not None:
            agent_status["structure_shadow"] = structure_shadow.summary
    elif source_backend == "legacy":
        source_bundle = prepare_legacy_runtime_source(
            file_path,
            resolved_sheet_name,
            source_format=source_format,
            header_rows=header_rows,
            anchor_detector=detect_anchor,
        )
        agent_status["source_ingestion"] = "legacy — authoritative parser"
    else:
        source_bundle = prepare_gated_runtime_source(
            file_path,
            resolved_sheet_name,
            source_format=source_format,
            header_rows=header_rows,
            llm_call=structure_llm_call,
            anchor_detector=detect_anchor,
        )
        structure_shadow = source_bundle.migration_report
        agent_status["source_ingestion"] = (
            "source-ir-gated — promoted after MATCH parity"
        )
        if structure_shadow is not None:
            agent_status["structure_shadow"] = structure_shadow.summary

    position_to_variety = source_bundle.position_to_variety
    variety_names_seen = source_bundle.variety_names
    attributes = source_bundle.schema_attributes
    on_progress(f"Header sumber: {[item.attribute_name for item in source_bundle.all_attributes]}")
    if source_bundle.anchor_attribute_name is not None:
        on_progress(f"Deteksi anchor: status='found' kolom={source_bundle.anchor_attribute_name!r}")
    if not variety_names_seen:
        raise ValueError("Tidak ada varietas untuk keluaran. Periksa header dan isi data sumber.")
    on_progress(f"Atribut untuk schema matching: {len(attributes)}")
    on_progress(f"Varietas terdeteksi dari sumber: {variety_names_seen}")

    # Optional observation-only shadow for the legacy backend. Its output
    # cannot affect downstream work. In source-ir-gated mode above, verified
    # Source IR is authoritative only after exact MATCH parity.
    if enable_structure_shadow and source_backend == "legacy":
        on_progress("structure_shadow: membandingkan legacy parser dengan Source IR...")
        try:
            structure_shadow = run_structure_shadow(
                file_path,
                resolved_sheet_name,
                source_format=source_format,
                header_rows=header_rows,
                llm_call=structure_llm_call,
                anchor_detector=detect_anchor,
            )
        except Exception as exc:  # noqa: BLE001 - shadow must never abort primary work
            message = sanitize_shadow_error_message(exc)
            structure_shadow = ShadowParityReport(
                status=ShadowStatus.NEW_PATH_FAILED,
                source_format=source_format,
                issue_codes=["SHADOW_RUNNER_EXCEPTION"],
                new_path_error_type=type(exc).__name__,
                new_path_error_message=message,
                summary="NEW_PATH_FAILED — isolated shadow exception; legacy pipeline continued",
            )
        agent_status["structure_shadow"] = structure_shadow.summary
        on_progress(f"structure_shadow: {structure_shadow.summary}")

    schema = CanonicalSchema.from_template()
    evaluation_config = EvaluationRunConfig(
        source_backend=source_backend,
        retrieval_backend=retrieval_backend,
        retrieval_k=k,
        canonical_schema_version=schema.schema_version,
        canonical_template_hash=schema.template_hash,
        mapping_verification_version=MAPPING_VERIFICATION_VERSION,
        embedding_model_name=EMBEDDING_MODEL_NAME,
    )
    retrieval_resource = None
    retrieval_initialized = False

    builder = CanonicalOutputBuilder(schema=schema)
    for name in variety_names_seen:
        builder.add_variety(name)

    # --- schema matching (Fase 3, via the Fase 7 reliability wrapper) ---
    mapping_rows: list[dict] = []
    provenance_records: list[CellProvenanceRecord] = []
    mapping_verifications: list[MappingVerificationResult] = []
    n_auto_accept = 0
    n_review = 0
    n_no_write = 0
    n_exact_name = 0
    n_retrieve_rerank = 0
    verifier_counts = {status: 0 for status in MappingVerificationStatus}
    n_hard_blocked = 0

    mapping_identities = build_mapping_item_identities(
        source_file_sha256=source_hash,
        source_sheet=resolved_sheet_name,
        source_format=source_format,
        source_items=[(attr.source_attribute_id, attr.display_name) for attr in attributes],
    )

    for attr, mapping_identity in zip(attributes, mapping_identities):
        mapping_item_id = mapping_identity.mapping_item_id
        on_progress(f"  schema_matching: '{attr.attribute_name}' — retrieval...")
        profile = SourceAttributeProfile(
            attribute_name=attr.attribute_name,
            structural_context=attr.structural_context,
            sample_values=attr.sample_values,
            header_path=attr.header_path,
            source_value_type=attr.detected_value_type,
            source_attribute_id=attr.source_attribute_id,
        )
        exact_resolution = resolve_exact_name(attr.attribute_name, schema)
        mapping_method: MappingMethod
        retrieved = None
        if exact_resolution.status is ExactNameStatus.MATCH:
            mapping_method = "exact_name"
            n_exact_name += 1
            mapping = mapping_from_exact_resolution(
                exact_resolution,
                source_attribute=profile.attribute_name,
                source_context=profile.structural_context,
                source_format=source_format,
            )
            patch = {}
            on_progress(
                f"  schema_matching: '{attr.attribute_name}' — exact_name; retrieval dilewati."
            )
        else:
            mapping_method = "retrieve_rerank"
            n_retrieve_rerank += 1
            if not retrieval_initialized:
                if retrieval_backend == "chroma":
                    on_progress("Memastikan indeks ChromaDB (idempoten)...")
                    retrieval_resource = ensure_indexed(schema)
                    agent_status["retrieval"] = f"chroma — HNSW cosine, k={k}"
                elif retrieval_backend == "exact":
                    on_progress("Menyiapkan indeks exact cosine dalam memori...")
                    build_kwargs = {}
                    if embedding_encode_call is not None:
                        build_kwargs["encode_call"] = embedding_encode_call
                    retrieval_resource = build_exact_index(schema, **build_kwargs)
                    agent_status["retrieval"] = (
                        f"exact — exhaustive cosine over {len(schema.rows)} canonical rows, k={k}"
                    )
                retrieval_initialized = True

            retrieval_kwargs = {"backend": retrieval_backend}
            if retrieval_backend == "chroma":
                retrieval_kwargs["collection"] = retrieval_resource
            elif retrieval_backend == "exact":
                retrieval_kwargs["exact_index"] = retrieval_resource
            if embedding_encode_call is not None:
                retrieval_kwargs["encode_call"] = embedding_encode_call
            retrieved = retrieve(profile, k=k, schema=schema, **retrieval_kwargs)

            on_progress(f"  schema_matching: '{attr.attribute_name}' — reranking...")
            mapping, patch = safe_rerank(
                profile, retrieved, state, source_format=source_format, schema=schema,
                enqueue_review=False, rate_limiter=text_rate_limiter,
            )
        state.update(patch)  # safe_* returns a patch; the caller applies it — see wrappers.py
        verification = verify_mapping(
            profile=profile,
            mapping=mapping,
            mapping_method=mapping_method,
            schema=schema,
            exact_resolution=exact_resolution,
            candidates=retrieved,
            source_format=source_format,
            reliability_patch=patch,
            mapping_item_id=mapping_item_id,
        )
        mapping_verifications.append(verification)
        verifier_counts[verification.status] += 1
        current_acceptance = decide_mapping_acceptance(mapping, reliability_patch=patch)
        acceptance = combine_mapping_acceptance(current_acceptance, verification)
        if verification.status is MappingVerificationStatus.REJECT:
            n_hard_blocked += 1

        target_row = (
            schema.row_by_id(mapping.target_canonical_row)
            if mapping is not None and mapping.target_canonical_row != NULL_ROW
            else None
        )
        mapping_row = {
            "mapping_item_id": mapping_item_id,
            "review_item_id": None,
            "mapping_identity_kind": mapping_identity.identity_kind.value,
            "mapping_identity_value": mapping_identity.identity_value,
            "mapping_identity_issue": mapping_identity.issue_code,
            "source_file_name": file_path.name,
            "source_file_sha256": source_hash,
            "source_sheet": resolved_sheet_name,
            "source_format": source_format,
            "source_attribute_display": attr.display_name,
            "source_attribute": attr.attribute_name,
            "source_context": attr.structural_context,
            "predicted_row": mapping.target_canonical_row if mapping is not None else None,
            "proposed_target_canonical_key": target_row.canonical_key if target_row else None,
            "predicted_label": target_row.label if target_row else None,
            "target_domain": mapping.target_domain if mapping is not None else None,
            "confidence": mapping.confidence if mapping is not None else None,
            "normalization_required": mapping.normalization_required if mapping is not None else None,
            "reasoning": mapping.reasoning if mapping is not None else None,
            "mapping_method": mapping_method,
            "exact_name_status": exact_resolution.status.value,
            "exact_name_candidates": list(exact_resolution.candidate_canonical_keys),
            "verifier_status": verification.status.value,
            "verifier_hard_issues": list(verification.hard_issue_codes),
            "verifier_warnings": list(verification.warning_codes),
            "retrieval_target_rank": (
                verification.retrieval_evidence.target_rank
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_target_distance": (
                verification.retrieval_evidence.target_distance
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_top1_row": (
                verification.retrieval_evidence.top1_row_id
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_top1_distance": (
                verification.retrieval_evidence.top1_distance
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_top2_row": (
                verification.retrieval_evidence.top2_row_id
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_top2_distance": (
                verification.retrieval_evidence.top2_distance
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_top1_top2_margin": (
                verification.retrieval_evidence.top1_top2_margin
                if verification.retrieval_evidence is not None
                else None
            ),
            "retrieval_target_vs_top1_gap": (
                verification.retrieval_evidence.target_vs_top1_distance_gap
                if verification.retrieval_evidence is not None
                else None
            ),
            "acceptance_status": acceptance.status.value,
            "acceptance_reason": acceptance.reason,
            "canonical_write": False,
        }
        mapping_rows.append(mapping_row)

        grouped = group_attribute_contributions_by_variety(attr, position_to_variety)
        if acceptance.status is AcceptanceStatus.REVIEW and mapping is not None:
            review_item = review_queue.enqueue(
                mapping,
                reason=acceptance.reason,
                queue_path=review_queue_path,
                run_id=run_id,
                mapping_item_id=mapping_item_id,
                source_file_name=file_path.name,
                source_file_sha256=source_hash,
                source_sheet=resolved_sheet_name,
                source_format=source_format,
                source_attribute_display=attr.display_name,
                source_attribute_id=attr.source_attribute_id,
                mapping_identity_kind=mapping_identity.identity_kind.value,
                mapping_identity_value=mapping_identity.identity_value,
                schema_version=schema.schema_version,
                template_hash=schema.template_hash,
                proposed_canonical_key=target_row.canonical_key if target_row else None,
                mapping_method=mapping_method,
                verifier_status=verification.status.value,
                verifier_warnings=list(verification.warning_codes),
                verifier_hard_issues=list(verification.hard_issue_codes),
                sample_values=list(attr.sample_values[:10]),
                raw_values_by_variety={
                    variety: [item.raw_value for item in contributions]
                    for variety, contributions in grouped.items()
                },
                source_cells_by_variety={
                    variety: physical_source_cells(contributions)
                    for variety, contributions in grouped.items()
                },
                source_header_cells=list(attr.header_cells),
                source_ir_version=(source_bundle.source_ir.ir_version if source_bundle.source_ir else None),
            )
            mapping_row["review_item_id"] = review_item.item_id
            state.update(review_queue.append_error_trace(state, acceptance.reason))

        # Selective-acceptance safety invariant: only AUTO_ACCEPT may cross
        # this boundary into normalization or canonical mutation.
        if not acceptance.allows_canonical_write:
            if acceptance.status is AcceptanceStatus.REVIEW:
                n_review += 1
            else:
                n_no_write += 1
            on_progress(
                f"  schema_matching: '{attr.attribute_name}' -> {acceptance.status.value}, "
                f"tidak ditulis: {acceptance.reason}"
            )
            continue

        n_auto_accept += 1

        on_progress(
            f"  schema_matching: '{attr.attribute_name}' -> {mapping.target_canonical_row} "
            f"(confidence={mapping.confidence:.2f}, AUTO_ACCEPT)"
        )

        if target_row is None:
            # Defensive fail-closed guard. The decision function must never
            # AUTO_ACCEPT an absent/NULL/unknown target.
            continue

        for variety_name, contributions in grouped.items():
            raw_values = [item.raw_value for item in contributions]
            combined = combine_multi_value(raw_values)
            if combined is None:
                continue
            normalized = normalize(combined, target_row)
            written = builder.set_cell(target_row.id, variety_name, normalized.value)
            if written:
                if normalized.note:
                    warning = (
                        f"normalization_warning: atribut={attr.attribute_name!r}, "
                        f"varietas={variety_name!r}: {normalized.note}"
                    )
                    state.update(review_queue.append_error_trace(state, warning))
                    on_progress(f"  schema_matching: {warning}")
                mapping_row["canonical_write"] = True
                provenance_records.append(
                    CellProvenanceRecord(
                        run_id=run_id,
                        source_file_name=file_path.name,
                        source_file_sha256=source_hash,
                        source_sheet=resolved_sheet_name,
                        source_attribute=attr.attribute_name,
                        source_context=attr.structural_context,
                        source_attribute_display=attr.display_name,
                        source_cells=physical_source_cells(contributions),
                        source_attribute_id=attr.source_attribute_id,
                        source_header_cells=list(attr.header_cells),
                        source_ir_version=(
                            source_bundle.source_ir.ir_version
                            if source_bundle.source_ir is not None
                            else None
                        ),
                        variety=variety_name,
                        canonical_row_id=target_row.id,
                        canonical_key=target_row.canonical_key,
                        canonical_label=target_row.label,
                        canonical_domain=target_row.domain,
                        raw_value=combined,
                        normalized_value=normalized.value,
                        normalization_required=mapping.normalization_required,
                        normalization_note=normalized.note,
                        mapping_confidence=mapping.confidence,
                        acceptance_status=acceptance.status.value,
                        acceptance_reason=acceptance.reason,
                        schema_version=schema.schema_version,
                        template_hash=schema.template_hash,
                        mapping_method=mapping_method,
                        verifier_status=verification.status.value,
                        verifier_hard_issues=list(verification.hard_issue_codes),
                    )
                )

    agent_status["schema_matching"] = (
        f"selesai — {len(mapping_rows)} atribut: {n_auto_accept} AUTO_ACCEPT, "
        f"{n_review} REVIEW, {n_no_write} NO_WRITE; methods: "
        f"{n_exact_name} exact_name, {n_retrieve_rerank} retrieve_rerank"
    )
    if not retrieval_initialized:
        agent_status["retrieval"] = (
            f"{retrieval_backend} — tidak diinisialisasi; semua atribut exact_name, k={k}"
        )
    agent_status["mapping_verifier"] = (
        f"verifier PASS={verifier_counts[MappingVerificationStatus.PASS]}, "
        f"REVIEW={verifier_counts[MappingVerificationStatus.REVIEW]}, "
        f"REJECT={verifier_counts[MappingVerificationStatus.REJECT]}; "
        f"hard-blocked={n_hard_blocked}"
    )

    mapping_df = pd.DataFrame(mapping_rows, columns=MAPPING_COLUMNS)
    if not mapping_df.empty:
        mapping_df = mapping_df.sort_values("confidence", ascending=True, kind="stable").reset_index(drop=True)

    # Materialize the tabular canonical workbook. Optional Drive discovery and
    # vision mutation are owned exclusively by the LangGraph runtime nodes.
    workbook = builder.build_workbook()
    worksheet = workbook[SHEET_NAME]

    canonical_df = worksheet_to_dataframe(worksheet, schema, builder.variety_names)

    workbook_bytes = _deterministic_workbook_bytes(workbook)

    thread_id: str | None = run_id
    agent_status["orchestrator"] = "dikelola LangGraph runtime"
    on_progress("Selesai.")
    return PipelineRunResult(
        mapping_df=mapping_df, canonical_df=canonical_df,
        workbook_bytes=workbook_bytes, vision_rows=[],
        provenance_records=provenance_records, agent_status=agent_status,
        checkpoint_thread_id=thread_id,
        error_trace=list(state.get("error_trace", [])),
        structure_shadow=structure_shadow, source_backend=source_backend,
        retrieval_backend=retrieval_backend,
        mapping_verifications=mapping_verifications,
        schema_version=schema.schema_version, template_hash=schema.template_hash,
        retrieval_k=k, mapping_verification_version=MAPPING_VERIFICATION_VERSION,
        embedding_model_name=EMBEDDING_MODEL_NAME,
        evaluation_config_fingerprint=evaluation_config.fingerprint,
        run_id=run_id, review_queue_path=str(review_queue_path),
        source_ir_version=(source_bundle.source_ir.ir_version if source_bundle.source_ir else None),
    )


def _first_sheet(path: Path) -> str:
    wb = openpyxl.load_workbook(path, read_only=True)
    name = wb.sheetnames[0]
    wb.close()
    return name


def _runtime_initial_state(
    file_path: Path, *, run_id: str, source_format: str, sheet_name: str | None,
    header_rows: int | None, drive_folder_id: str | None, k: int, max_images: int,
    enable_structure_shadow: bool, source_backend: str, retrieval_backend: str,
    review_queue_path: Path | str,
) -> dict:
    from src.orchestrator.graph import runtime_identity

    if source_backend not in {"legacy", "source-ir-gated"}:
        raise ValueError(f"unknown source backend {source_backend!r}; expected one of: legacy, source-ir-gated")
    validate_retrieval_backend(retrieval_backend)
    schema = CanonicalSchema.from_template()
    resolved_sheet = sheet_name or _first_sheet(file_path)
    evaluation = EvaluationRunConfig(
        source_backend=source_backend, retrieval_backend=retrieval_backend, retrieval_k=k,
        canonical_schema_version=schema.schema_version,
        canonical_template_hash=schema.template_hash,
        mapping_verification_version=MAPPING_VERIFICATION_VERSION,
        embedding_model_name=EMBEDDING_MODEL_NAME,
    )
    state = {
        "run_id": run_id, "source_path": str(file_path), "source_file_name": file_path.name,
        "source_file_sha256": source_file_sha256(file_path), "source_format": source_format,
        "sheet_name": resolved_sheet, "header_rows": header_rows,
        "drive_folder_id": (drive_folder_id or "").strip(), "max_images": max_images,
        "enable_structure_shadow": enable_structure_shadow,
        "source_backend": source_backend, "retrieval_backend": retrieval_backend,
        "retrieval_k": k, "schema_version": schema.schema_version,
        "template_hash": schema.template_hash,
        "evaluation_config_fingerprint": evaluation.fingerprint,
        "mapping_verification_version": MAPPING_VERIFICATION_VERSION,
        "embedding_model_name": EMBEDDING_MODEL_NAME,
        "review_queue_path": str(review_queue_path), "error_trace": [], "agent_status": {},
    }
    state["runtime_identity"] = runtime_identity(state)
    return state


def _result_from_runtime_state(state: dict) -> PipelineRunResult:
    from src.schema.shadow_parity import ShadowParityReport

    mapping_df = pd.read_json(io.StringIO(state["mapping_json"]), orient="table")
    schema = CanonicalSchema.from_template()
    workbook = openpyxl.load_workbook(io.BytesIO(state["workbook_bytes"]))
    canonical_df = worksheet_to_dataframe(workbook[SHEET_NAME], schema, state["variety_names"])
    workbook.close()
    return PipelineRunResult(
        mapping_df=mapping_df, canonical_df=canonical_df,
        workbook_bytes=state["workbook_bytes"], vision_rows=list(state.get("vision_rows", [])),
        images_discovered=len(state.get("image_metadata", [])),
        provenance_records=[CellProvenanceRecord.model_validate(x) for x in state.get("provenance_records", [])],
        agent_status=dict(state.get("agent_status", {})), checkpoint_thread_id=state["run_id"],
        error_trace=list(state.get("error_trace", [])),
        structure_shadow=(ShadowParityReport.model_validate(state["structure_shadow"]) if state.get("structure_shadow") else None),
        source_backend=state["source_backend"], retrieval_backend=state["retrieval_backend"],
        mapping_verifications=[MappingVerificationResult.model_validate(x) for x in state.get("mapping_verifications", [])],
        schema_version=state["schema_version"], template_hash=state["template_hash"],
        retrieval_k=state["retrieval_k"], mapping_verification_version=state["mapping_verification_version"],
        embedding_model_name=state["embedding_model_name"],
        evaluation_config_fingerprint=state["evaluation_config_fingerprint"], run_id=state["run_id"],
        review_queue_path=state["review_queue_path"], source_ir_version=state.get("source_ir_version"),
    )


def _run_pipeline_ui_impl(
    file_path: Path, *, source_format: str = "row-oriented", sheet_name: str | None = None,
    header_rows: int | None = None, drive_folder_id: str | None = None, k: int = DEFAULT_K,
    max_images: int = 5, on_progress: ProgressCallback = _noop,
    enable_structure_shadow: bool = False, structure_llm_call: Callable | None = None,
    source_backend: Literal["legacy", "source-ir-gated"] = "legacy",
    retrieval_backend: RetrievalBackend = "chroma", embedding_encode_call: Callable[..., object] | None = None,
    review_queue_path: Path | str = review_queue.DEFAULT_QUEUE_PATH,
    text_rate_limiter: RateLimiter | None = None, vision_rate_limiter: RateLimiter | None = None,
    checkpoint_db_path: Path | str | None = None, _run_id: str | None = None,
    _interrupt_after: list[str] | None = None,
) -> PipelineRunResult | dict:
    from src.orchestrator.graph import DEFAULT_CHECKPOINT_DB, RuntimeResources, run_pipeline

    run_id = _run_id or uuid.uuid4().hex
    initial = _runtime_initial_state(file_path, run_id=run_id, source_format=source_format,
        sheet_name=sheet_name, header_rows=header_rows, drive_folder_id=drive_folder_id, k=k,
        max_images=max_images, enable_structure_shadow=enable_structure_shadow,
        source_backend=source_backend, retrieval_backend=retrieval_backend,
        review_queue_path=review_queue_path)
    initial["agent_status"] = {
        "text_rate_limit": describe_rate_limiter(text_rate_limiter),
        "vision_rate_limit": describe_rate_limiter(vision_rate_limiter),
    }
    resources = RuntimeResources(on_progress=on_progress, text_rate_limiter=text_rate_limiter,
        vision_rate_limiter=vision_rate_limiter, structure_llm_call=structure_llm_call,
        embedding_encode_call=embedding_encode_call)
    final = run_pipeline(initial, resources, db_path=checkpoint_db_path or DEFAULT_CHECKPOINT_DB,
                         interrupt_after=_interrupt_after)
    return _result_from_runtime_state(final) if final.get("completed") else final


def _resume_pipeline_ui_impl(
    file_path: Path, *, run_id: str, source_format: str = "row-oriented",
    sheet_name: str | None = None, header_rows: int | None = None,
    drive_folder_id: str | None = None, k: int = DEFAULT_K, max_images: int = 5,
    enable_structure_shadow: bool = False, structure_llm_call: Callable | None = None,
    source_backend: Literal["legacy", "source-ir-gated"] = "legacy",
    retrieval_backend: RetrievalBackend = "chroma", embedding_encode_call: Callable[..., object] | None = None,
    review_queue_path: Path | str = review_queue.DEFAULT_QUEUE_PATH,
    text_rate_limiter: RateLimiter | None = None, vision_rate_limiter: RateLimiter | None = None,
    checkpoint_db_path: Path | str | None = None, on_progress: ProgressCallback = _noop,
) -> PipelineRunResult:
    from src.orchestrator.graph import DEFAULT_CHECKPOINT_DB, RuntimeResources, resume_pipeline

    expected = _runtime_initial_state(file_path, run_id=run_id, source_format=source_format,
        sheet_name=sheet_name, header_rows=header_rows, drive_folder_id=drive_folder_id, k=k,
        max_images=max_images, enable_structure_shadow=enable_structure_shadow,
        source_backend=source_backend, retrieval_backend=retrieval_backend,
        review_queue_path=review_queue_path)
    resources = RuntimeResources(on_progress=on_progress, text_rate_limiter=text_rate_limiter,
        vision_rate_limiter=vision_rate_limiter, structure_llm_call=structure_llm_call,
        embedding_encode_call=embedding_encode_call)
    final = resume_pipeline(expected, resources, db_path=checkpoint_db_path or DEFAULT_CHECKPOINT_DB)
    return _result_from_runtime_state(final)


def run_pipeline_ui(
    file_path: Path,
    *,
    source_format: str = "row-oriented",
    sheet_name: str | None = None,
    header_rows: int | None = None,
    drive_folder_id: str | None = None,
    k: int = DEFAULT_K,
    max_images: int = 5,
    on_progress: ProgressCallback = _noop,
    enable_structure_shadow: bool = False,
    structure_llm_call: Callable | None = None,
    source_backend: Literal["legacy", "source-ir-gated"] = "legacy",
    retrieval_backend: RetrievalBackend = "chroma",
    embedding_encode_call: Callable[..., object] | None = None,
    review_queue_path: Path | str = review_queue.DEFAULT_QUEUE_PATH,
    text_rate_limiter: RateLimiter | None = None,
    vision_rate_limiter: RateLimiter | None = None,
    rate_limit_config: RuntimeRateLimitConfig | None = None,
    checkpoint_db_path: Path | str | None = None,
) -> PipelineRunResult:
    """Run the UI pipeline with per-run text and vision limiter ownership.

    Limiters supplied by a caller are borrowed and never closed here. A
    missing limiter is created from the explicit config (or environment) and
    is always closed by this function, including when the pipeline raises.
    """
    config = rate_limit_config or RuntimeRateLimitConfig.from_env()
    owned_limiters: list[RateLimiter] = []
    try:
        if text_rate_limiter is None and config.text_rpm is not None:
            text_rate_limiter = RateLimiter(config.text_rpm, 60.0)
            owned_limiters.append(text_rate_limiter)
        if vision_rate_limiter is None and config.vision_rpm is not None:
            vision_rate_limiter = RateLimiter(config.vision_rpm, 60.0)
            owned_limiters.append(vision_rate_limiter)
        return _run_pipeline_ui_impl(
            file_path,
            source_format=source_format,
            sheet_name=sheet_name,
            header_rows=header_rows,
            drive_folder_id=drive_folder_id,
            k=k,
            max_images=max_images,
            on_progress=on_progress,
            enable_structure_shadow=enable_structure_shadow,
            structure_llm_call=structure_llm_call,
            source_backend=source_backend,
            retrieval_backend=retrieval_backend,
            embedding_encode_call=embedding_encode_call,
            review_queue_path=review_queue_path,
            text_rate_limiter=text_rate_limiter,
            vision_rate_limiter=vision_rate_limiter,
            checkpoint_db_path=checkpoint_db_path,
        )
    finally:
        for limiter in reversed(owned_limiters):
            limiter.close()


def resume_pipeline_ui(
    file_path: Path, *, run_id: str, source_format: str = "row-oriented",
    sheet_name: str | None = None, header_rows: int | None = None,
    drive_folder_id: str | None = None, k: int = DEFAULT_K, max_images: int = 5,
    enable_structure_shadow: bool = False, structure_llm_call: Callable | None = None,
    source_backend: Literal["legacy", "source-ir-gated"] = "legacy",
    retrieval_backend: RetrievalBackend = "chroma", embedding_encode_call: Callable[..., object] | None = None,
    review_queue_path: Path | str = review_queue.DEFAULT_QUEUE_PATH,
    text_rate_limiter: RateLimiter | None = None, vision_rate_limiter: RateLimiter | None = None,
    rate_limit_config: RuntimeRateLimitConfig | None = None,
    checkpoint_db_path: Path | str | None = None, on_progress: ProgressCallback = _noop,
) -> PipelineRunResult:
    """Resume a checkpointed graph with fresh non-checkpointed runtime resources."""
    config = rate_limit_config or RuntimeRateLimitConfig.from_env()
    owned: list[RateLimiter] = []
    try:
        if text_rate_limiter is None and config.text_rpm is not None:
            text_rate_limiter = RateLimiter(config.text_rpm, 60.0); owned.append(text_rate_limiter)
        if vision_rate_limiter is None and config.vision_rpm is not None:
            vision_rate_limiter = RateLimiter(config.vision_rpm, 60.0); owned.append(vision_rate_limiter)
        return _resume_pipeline_ui_impl(file_path, run_id=run_id, source_format=source_format,
            sheet_name=sheet_name, header_rows=header_rows, drive_folder_id=drive_folder_id,
            k=k, max_images=max_images, enable_structure_shadow=enable_structure_shadow,
            structure_llm_call=structure_llm_call, source_backend=source_backend,
            retrieval_backend=retrieval_backend, embedding_encode_call=embedding_encode_call,
            review_queue_path=review_queue_path, text_rate_limiter=text_rate_limiter,
            vision_rate_limiter=vision_rate_limiter, checkpoint_db_path=checkpoint_db_path,
            on_progress=on_progress)
    finally:
        for limiter in reversed(owned): limiter.close()
