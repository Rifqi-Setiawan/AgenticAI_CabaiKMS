from __future__ import annotations

import hashlib
import io
import types
from typing import Any

import openpyxl

from .client import KmsClient
from src.agents import vision_classification
from src.agents.signal_fusion import FusionAction, fuse_photo_signals
from src.agents.variety_registry import VarietyRegistry
from src.schema.contracts import ImageMetadata
from src.storage.config import get_settings
from src.storage.errors import InvalidState
from src.ui import pipeline_runner as runner


def extract_run_scope_varieties(workbook: openpyxl.Workbook) -> list[str]:
    """Extract variety names detected in the active Excel run."""
    varieties: list[str] = []
    # Check sheets for variety column
    for sheet in workbook.worksheets:
        for r in range(1, min(5, sheet.max_row + 1)):
            for c in range(1, sheet.max_column + 1):
                val = str(sheet.cell(row=r, column=c).value or "").strip().lower()
                if val in {"jenis cabai", "varietas", "nama varietas"}:
                    for data_row in range(r + 1, sheet.max_row + 1):
                        v_val = sheet.cell(row=data_row, column=c).value
                        if v_val is not None:
                            v_str = str(v_val).strip()
                            if v_str and v_str not in varieties:
                                varieties.append(v_str)
                    if varieties:
                        return varieties
    # Check if header row has varieties across columns (e.g. canonical template)
    if "Sheet1" in workbook.sheetnames:
        ws = workbook["Sheet1"]
        row1 = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        if len(row1) > 2 and row1[1] and "karakteristik" in str(row1[1]).lower():
            for v in row1[2:]:
                if v and str(v).strip() and str(v).strip() not in varieties:
                    varieties.append(str(v).strip())
    return varieties


def vision_classification_with_storage(state: dict[str, Any], config: dict[str, Any]) -> dict[str, Any]:
    """Persist first; inference consumes the exact scrubbed derivative recorded in evidence."""
    from src.orchestrator.graph import _resources

    resources = _resources(config)
    client = KmsClient(get_settings())
    workbook = openpyxl.load_workbook(io.BytesIO(state['workbook_bytes']))
    try:
        if runner.SHEET_NAME in workbook.sheetnames:
            worksheet = workbook[runner.SHEET_NAME]
        elif "Resume Data" in workbook.sheetnames:
            worksheet = workbook["Resume Data"]
        else:
            worksheet = workbook.active

        run_scope_varieties = extract_run_scope_varieties(workbook)
        db_varieties = []
        try:
            var_res = client.request('GET', '/varieties')
            items = var_res.get('items', var_res) if isinstance(var_res, dict) else var_res
            for item in (items or []):
                if item:
                    db_varieties.append(item)
        except Exception as exc:
            resources.on_progress(f"vision: Gagal mengambil varietas DB: {exc}")
            raise RuntimeError(f"Failed to fetch DB varieties: {exc}") from exc

        registry = VarietyRegistry.build(spreadsheet_varieties=run_scope_varieties, db_varieties=db_varieties)
        # F-6: Combined variety registry options provided to VisionSession
        all_eval_options = registry.get_canonical_names()
        vision = vision_classification.VisionSession(run_scope_varieties=all_eval_options)

        trace = {
            'error_trace': list(state.get('error_trace', [])),
            'provider_trace': list(state.get('provider_trace', [])),
        }
        rows: list[dict[str, Any]] = []
        written = 0
        review_items = list(state.get('review_items', []))
        total_images = len(state.get('image_metadata', []))

        for idx, raw in enumerate(state.get('image_metadata', []), start=1):
            image = ImageMetadata.model_validate(raw)
            resources.on_progress(f"vision [{idx}/{total_images}]: Mengunduh '{image.filename}' ke CAS storage...")
            # Failure is visible as a failed node; never silently persist a DB reference to a missing file.
            asset = client.request('POST', '/assets/drive/' + image.file_id)

            # Check if asset original_sha256 was previously discarded
            is_discarded = False
            discard_reason = "Image previously discarded by curator."
            try:
                disc_res = client.request('GET', f"/discarded?sha256={asset.get('original_sha256', '')}")
                if disc_res and isinstance(disc_res, list) and len(disc_res) > 0:
                    is_discarded = True
                    discard_reason = disc_res[0].get('reason') or discard_reason
            except Exception as exc:
                resources.on_progress(f"vision: Gagal memeriksa discarded images: {exc}")
                raise RuntimeError(f"Failed to check discarded status for {image.filename}: {exc}") from exc

            if is_discarded:
                resources.on_progress(f"vision [{idx}/{total_images}]: Melewati '{image.filename}' (telah ditolak/discarded)")
                rows.append({
                    'filename': image.filename,
                    'file_id': image.file_id,
                    'asset_id': asset['id'],
                    'status': 'DISCARDED',
                    'write_applied': False,
                    'write_reason': f"Discarded image: {discard_reason}",
                    'options_evaluated': run_scope_varieties,
                    'signals_consistent': False,
                    'original_sha256': asset['original_sha256'],
                })
                continue

            image_bytes = client.request('GET', f"/media/{asset['id']}/image")
            actual_sha = hashlib.sha256(image_bytes).hexdigest()
            expected_sha = asset.get('image_sha256')
            if expected_sha and len(expected_sha) == 64 and actual_sha != expected_sha:
                raise InvalidState('Inference media checksum mismatch.')

            model_image = image.model_copy(update={'mime_type': 'image/webp'})
            resources.on_progress(f"vision [{idx}/{total_images}]: Menganalisis '{image.filename}' via model multimodal...")

            photo_signal = vision_classification.extract_photo_context_signal(
                image.filename,
                known_varieties=run_scope_varieties,
                registry=registry,
            )

            try:
                result = vision.classify(
                    model_image,
                    image_bytes=image_bytes,
                    photo_signal=photo_signal,
                )
            except Exception as exc:
                trace['error_trace'].append(f"Vision classification failed for {image.filename}: {exc}")
                result = None

            if result is None:
                resources.on_progress(f"vision [{idx}/{total_images}]: Gagal klasifikasi '{image.filename}'")
                rows.append({
                    'filename': image.filename,
                    'file_id': image.file_id,
                    'asset_id': asset['id'],
                    'status': 'FAILED',
                    'write_applied': False,
                    'write_reason': 'Vision provider/contract failed.',
                    'options_evaluated': run_scope_varieties,
                    'signals_consistent': False,
                })
                continue

            # Signal fusion logic
            outcome = fuse_photo_signals(
                filename_match=photo_signal.suggested_variety_hint,
                registry_known=registry,
                visual_result=result,
            )

            review_item = None
            if outcome.action == FusionAction.AUTO_LINK:
                try:
                    update = runner.apply_vision_result_to_worksheet(worksheet, image, result)
                except Exception as exc:
                    update = types.SimpleNamespace(applied=False, reason=str(exc))
                status_to_report = result.classification_status
            elif outcome.action == FusionAction.CANDIDATE:
                update = types.SimpleNamespace(
                    applied=False,
                    reason=f"CANDIDATE: {outcome.reason}",
                )
                status_to_report = "CANDIDATE"
                cand_payload = {
                    "asset_id": asset['id'],
                    "run_id": state.get('run_id') or "current_run",
                    "source_filename": image.filename,
                    "original_sha256": asset['original_sha256'],
                    "organ_part": (result.identified_part or "buah").lower(),
                    "confidence": float(getattr(result, "visual_confidence", getattr(result, "confidence", 0.0))),
                    "suggested_name": outcome.suggested_candidate_name,
                    "visual_evidence": {"evidence": result.visual_evidence},
                }
                try:
                    client.request('POST', '/candidates', json=cand_payload)
                except Exception as e:
                    trace['error_trace'].append(f"Failed to record candidate {image.filename}: {e}")
                    status_to_report = "FAILED_CANDIDATE_RECORD"
                    update.reason = f"FAILED_CANDIDATE_RECORD: {e}"
            else:  # REVIEW_MISMATCH or REVIEW_UNCERTAIN
                review_item = vision.evaluate_photo_alignment(
                    image,
                    photo_signal=photo_signal,
                    visual_result=result,
                )
                if review_item is None:
                    review_item = vision_classification.create_photo_mismatch_review_item(
                        filename=image.filename,
                        suggested_hint=photo_signal.suggested_variety_hint,
                        visual_evidence=result.visual_evidence,
                        options_evaluated=run_scope_varieties,
                        caption=photo_signal.caption,
                        identified_part=result.identified_part,
                        confidence=getattr(result, 'confidence', 0.0),
                        source_context=outcome.reason,
                    )
                review_items.append(review_item)
                update = types.SimpleNamespace(
                    applied=False,
                    reason=f"PHOTO_VARIETY_MISMATCH: {outcome.reason}",
                )
                status_to_report = result.classification_status

            written += int(update.applied)
            write_status = (
                f"ditulis ke sel '{result.identified_part}' ({result.matched_variety})"
                if update.applied
                else f"tidak ditulis ({update.reason})"
            )
            resources.on_progress(
                f"vision [{idx}/{total_images}]: '{image.filename}' -> "
                f"{status_to_report} [{result.identified_part}, {result.matched_variety}] ({write_status})"
            )

            rows.append({
                'filename': image.filename,
                'file_id': image.file_id,
                'status': status_to_report,
                'matched_variety': result.matched_variety,
                'identified_part': result.identified_part,
                'confidence': getattr(result, 'confidence', 0.0),
                'visual_evidence': result.visual_evidence,
                'write_applied': update.applied,
                'write_reason': update.reason,
                'asset_id': asset['id'],
                'original_sha256': asset['original_sha256'],
                'inference_sha256': asset['image_sha256'],
                'transform_profile': asset['transform_profile'],
                'options_evaluated': getattr(result, 'options_evaluated', run_scope_varieties),
                'signals_consistent': getattr(result, 'signals_consistent', review_item is None),
            })

        status = dict(state.get('agent_status', {}))
        status['vision_classification'] = f'{len(rows)} classifications; {written} writes; immutable media captured'
        resources.on_progress(f"vision_classification: selesai ({len(rows)} citra diproses, {written} tautan kanonik ditulis)")

        return {
            'workbook_bytes': runner._deterministic_workbook_bytes(workbook),
            'vision_rows': rows,
            'review_items': review_items,
            'error_trace': trace['error_trace'],
            'provider_trace': trace.get('provider_trace', []),
            'agent_status': status,
        }
    finally:
        workbook.close()
